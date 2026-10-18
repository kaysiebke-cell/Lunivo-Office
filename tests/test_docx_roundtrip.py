import io

import docx
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PySide6.QtGui import QImage, QColor, QTextCursor

from native_qt.document_io.docx_reader import read_docx
from native_qt.document_io.docx_writer import write_docx
from native_qt.document_io.report import DocumentIOError
import pytest


def _png():
    img = QImage(40, 20, QImage.Format_RGB32)
    img.fill(QColor("red"))
    from PySide6.QtCore import QBuffer, QIODevice
    b = QBuffer(); b.open(QIODevice.WriteOnly); img.save(b, "PNG")
    return io.BytesIO(bytes(b.data()))


@pytest.fixture
def sample(tmp_path):
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = Mm(30)
    d.add_heading("Titel Eins", 1)
    p = d.add_paragraph()
    r = p.add_run("fett "); r.bold = True
    r = p.add_run("kursiv "); r.italic = True
    r = p.add_run("rot"); r.font.color.rgb = RGBColor(255, 0, 0); r.font.size = Pt(18); r.font.name = "DejaVu Serif"
    q = d.add_paragraph("zentriert"); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.left_indent = Mm(10)
    d.add_paragraph("Punkt A", style="List Bullet")
    d.add_paragraph("Punkt B", style="List Bullet")
    d.add_paragraph("Eins", style="List Number")
    d.add_paragraph("Zwei", style="List Number")
    t = d.add_table(rows=2, cols=2); t.style = "Table Grid"
    t.cell(0, 0).text = "A1"; t.cell(0, 1).text = "B1"; t.cell(1, 0).text = "A2"; t.cell(1, 1).text = "B2"
    d.add_picture(_png(), width=Mm(30))
    d.add_page_break()
    d.add_paragraph("Seite zwei")
    s.footer.is_linked_to_previous = False
    s.footer.paragraphs[0].text = "Fuß"
    f = tmp_path / "in.docx"; d.save(f)
    return str(f)


def texts(path):
    return [p.text for p in docx.Document(path).paragraphs]


def test_read_basic(app, sample):
    qd, page, meta, rep = read_docx(sample)
    txt = qd.toPlainText()
    for w in ("Titel Eins", "fett kursiv rot", "zentriert", "Punkt A", "Zwei", "A1", "B2", "Seite zwei"):
        assert w in txt
    assert abs(page.margin_left_mm - 30) < 0.1
    assert page.footer.text == "Fuß"
    assert qd.firstBlock().blockFormat().headingLevel() == 1


def test_roundtrip_preserves_content_and_format(app, sample, tmp_path):
    qd, page, meta, _ = read_docx(sample)
    out = str(tmp_path / "out.docx")
    write_docx(out, qd, page, meta)
    d = docx.Document(out)
    ps = d.paragraphs
    assert ps[0].style.name == "Heading 1" and ps[0].text == "Titel Eins"
    runs = ps[1].runs
    assert runs[0].bold and "fett" in runs[0].text
    assert any(r.italic for r in runs)
    red = [r for r in runs if r.font.color.rgb == RGBColor(255, 0, 0)]
    assert red and red[0].font.size == Pt(18) and red[0].font.name == "DejaVu Serif"
    cen = [p for p in ps if p.text == "zentriert"][0]
    assert cen.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert abs(cen.paragraph_format.left_indent.mm - 10) < 0.5
    assert [p.style.name for p in ps if p.text.startswith("Punkt")] == ["List Bullet"] * 2
    assert [p.style.name for p in ps if p.text in ("Eins", "Zwei")] == ["List Number"] * 2
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "B2"
    assert len(d.inline_shapes) == 1
    assert any(p.paragraph_format.page_break_before for p in ps if p.text == "Seite zwei")
    assert abs(d.sections[0].left_margin.mm - 30) < 0.1
    assert d.sections[0].footer.paragraphs[0].text == "Fuß"


def test_second_roundtrip_is_stable(app, sample, tmp_path):
    qd, page, meta, _ = read_docx(sample)
    a = str(tmp_path / "a.docx"); write_docx(a, qd, page, meta)
    qd2, page2, meta2, _ = read_docx(a)
    b = str(tmp_path / "b.docx"); write_docx(b, qd2, page2, meta2)
    assert texts(a) == texts(b)
    assert qd2.toPlainText() == qd.toPlainText()


def test_numbering_restarts_per_list(app, tmp_path):
    from PySide6.QtGui import QTextDocument, QTextListFormat
    qd = QTextDocument(); c = QTextCursor(qd)
    lf = QTextListFormat(); lf.setStyle(QTextListFormat.ListDecimal)
    c.insertList(lf); c.insertText("a"); c.insertBlock(); c.insertText("b")
    c.insertBlock(); c.setBlockFormat(__import__("PySide6.QtGui", fromlist=["x"]).QTextBlockFormat())
    c.insertText("Zwischentext")
    lf2 = QTextListFormat(); lf2.setStyle(QTextListFormat.ListDecimal)
    c.insertList(lf2); c.insertText("c")
    out = str(tmp_path / "n.docx")
    from native_qt.editor.page import PageSettings
    write_docx(out, qd, PageSettings(), {})
    d = docx.Document(out)
    nums = {p.text: p._p.pPr.numPr.numId.val for p in d.paragraphs if p.style.name == "List Number"}
    assert nums["a"] == nums["b"] != nums["c"]


def test_invalid_file(app, tmp_path):
    bad = tmp_path / "kaputt.docx"; bad.write_bytes(b"kein zip")
    with pytest.raises(DocumentIOError):
        read_docx(str(bad))


def test_unsupported_content_is_reported(app, tmp_path):
    d = docx.Document(); p = d.add_paragraph("x")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = OxmlElement("w:r"); fr = OxmlElement("w:footnoteReference"); fr.set(qn("w:id"), "1"); r.append(fr); p._p.append(r)
    f = tmp_path / "f.docx"; d.save(f)
    _, _, _, rep = read_docx(str(f))
    assert any("Fußnoten" in w for w in rep.warnings)
