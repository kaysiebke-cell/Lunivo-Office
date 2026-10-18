import docx
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QFont, QTextCursor

from native_qt.document_io import formats
from native_qt.editor.editor_widget import DocumentEditor
from native_qt.editor.page import PageSettings


def make(app, text="Hallo Welt\nZweite Zeile"):
    e = DocumentEditor(); e.setPlainText(text); return e


def select_all(e):
    c = e.textCursor(); c.select(QTextCursor.Document); e.setTextCursor(c)


def test_format_commands_land_in_docx(app, tmp_path):
    e = make(app)
    select_all(e)
    e.toggle_bold(True); e.set_font_size(16); e.set_font_family("DejaVu Sans")
    e.set_text_color(QColor("#ff0000")); e.set_alignment(Qt.AlignHCenter); e.set_line_spacing(1.5)
    e.change_indent(10)
    out = str(tmp_path / "f.docx")
    formats.save(out, e.document(), e.page)
    d = docx.Document(out)
    p = d.paragraphs[0]
    assert p.alignment == 1
    assert p.runs[0].bold and p.runs[0].font.size.pt == 16 and p.runs[0].font.name == "DejaVu Sans"
    assert str(p.runs[0].font.color.rgb) == "FF0000"
    assert abs(p.paragraph_format.left_indent.mm - 10) < 0.5
    assert abs(p.paragraph_format.line_spacing - 1.5) < 0.01


def test_styles_and_lists(app, tmp_path):
    e = make(app, "Kopf\nA\nB")
    c = e.textCursor(); c.setPosition(0); e.setTextCursor(c)
    e.apply_style("Überschrift 2")
    assert e.current_style() == "Überschrift 2"
    c = e.textCursor(); c.movePosition(QTextCursor.NextBlock); c.movePosition(QTextCursor.NextBlock, QTextCursor.KeepAnchor); e.setTextCursor(c)
    e.toggle_list("bullet")
    out = str(tmp_path / "s.docx"); formats.save(out, e.document(), e.page)
    d = docx.Document(out)
    assert [p.style.name for p in d.paragraphs] == ["Heading 2", "List Bullet", "List Bullet"]


def test_find_replace_all(app):
    e = make(app, "foo Foo bar foo")
    assert e.replace_all("foo", "X") == 3 and e.toPlainText() == "X X bar X"
    e.undo()
    assert e.toPlainText() == "foo Foo bar foo"          # ein Schritt rückgängig
    assert e.replace_all("foo", "X", case=True) == 2


def test_find_wraps(app):
    e = make(app, "abc abc")
    assert e.find_text("abc") and e.find_text("abc") and e.find_text("abc")


def test_page_count_and_margin_change(app):
    e = make(app, "\n".join(f"Zeile {i}" for i in range(120)))
    e._recount()
    n = e.page_count
    assert n >= 2
    pg = e.page.copy(); pg.margin_left_mm = pg.margin_right_mm = 60
    e.page_settings_changed(pg); e._recount()
    assert e.page_count >= n and e.document().isModified()


def test_table_image_pagebreak_roundtrip(app, tmp_path):
    from PySide6.QtGui import QImage
    img = QImage(100, 50, QImage.Format_RGB32); img.fill(QColor("blue")); ip = str(tmp_path / "i.png"); img.save(ip)
    e = make(app, "Vor")
    e.moveCursor(QTextCursor.End)
    e.insert_table(2, 3); e.moveCursor(QTextCursor.End)
    assert e.insert_image(ip)
    e.insert_page_break(); e.insertPlainText("Nach")
    out = str(tmp_path / "t.docx"); formats.save(out, e.document(), e.page)
    d = docx.Document(out)
    assert len(d.tables) == 1 and len(d.tables[0].rows) == 2 and len(d.tables[0].columns) == 3
    assert len(d.inline_shapes) == 1
    assert any(p.paragraph_format.page_break_before for p in d.paragraphs)
    # und wieder einlesen
    l = formats.load(out)
    assert "Nach" in l.doc.toPlainText()


def test_seitenwechsel_werden_gefunden(app):
    e = make(app, "\n".join(f"Zeile {i}" for i in range(150)))
    e._recount()
    b = e.seitenwechsel
    assert len(b) == e.page_count - 1 >= 2
    assert b == sorted(b) and all(0 < x < e.document().characterCount() for x in b)
    # Seitenumbruch erzwingt einen Wechsel
    e2 = make(app, "A"); e2.moveCursor(QTextCursor.End); e2.insert_page_break(); e2.insertPlainText("B"); e2._recount()
    assert len(e2.seitenwechsel) == 1


def _tippe(e, text):
    from PySide6.QtGui import QKeyEvent
    from PySide6.QtCore import QEvent
    for ch in text:
        e.keyPressEvent(QKeyEvent(QEvent.KeyPress, Qt.Key_Space if ch == " " else 0, Qt.NoModifier, ch))


def test_autoformat_liste_und_url(app):
    e = make(app, "")
    e.opt = lambda n, std=True: True
    _tippe(e, "1. ")
    assert e.textCursor().currentList() is not None
    e2 = make(app, "")
    e2.opt = lambda n, std=True: True
    _tippe(e2, "- ")
    assert e2.textCursor().currentList() is not None


def test_seitenliste_duplex_umgekehrt():
    from native_qt.editor.pagination import seiten_liste
    assert seiten_liste(5, 1, 5, {"duplex": "vorder"}) == [1, 3, 5]
    assert seiten_liste(5, 1, 5, {"duplex": "rueck"}) == [2, 4]
    assert seiten_liste(3, 1, 3, {"umgekehrt": True}) == [3, 2, 1]


def test_design_und_absatzabstand(app):
    from native_qt.windows.befehle import Befehle
    e = make(app, "Titel\nText")
    c = e.textCursor(); c.setPosition(0); e.setTextCursor(c)
    e.apply_style("Überschrift 1")
    e.apply_design("Liberation Serif", "DejaVu Sans", "#2f6fb5")
    k = e.document().firstBlock().begin().fragment().charFormat()
    assert "DejaVu Sans" in k.fontFamilies() and k.foreground().color().name() == "#2f6fb5"
    t = e.document().firstBlock().next().begin().fragment().charFormat()
    assert "Liberation Serif" in t.fontFamilies() and t.foreground().color().name() != "#2f6fb5"
    e.select_range(0, e.document().characterCount() - 1)
    e.set_paragraph_spacing(6, 12)
    bf = e.document().firstBlock().next().blockFormat()
    assert abs(bf.bottomMargin() - 16) < 0.1


def test_zeilennummern_im_pdf(app, tmp_path):
    import pypdf
    from PySide6.QtGui import QTextDocument
    from native_qt.editor.page import PageSettings
    from native_qt.editor.pagination import export_pdf
    d = QTextDocument(); d.setPlainText("\n".join(f"Zeile{i}" for i in range(1, 8)))
    out = str(tmp_path / "n.pdf")
    export_pdf(out, d, PageSettings(zeilennummern=True), "t")
    txt = pypdf.PdfReader(out).pages[0].extract_text()
    assert all(str(i) in txt for i in range(1, 8))
    out2 = str(tmp_path / "m.pdf")
    export_pdf(out2, d, PageSettings(), "t")
    assert pypdf.PdfReader(out2).pages[0].extract_text().count("7") <= pypdf.PdfReader(out).pages[0].extract_text().count("7")


def test_formatvorlagen_verwalten_wirkt_auf_alle_absaetze(app):
    from native_qt.editor import stile
    e = make(app, "Eins\nZwei\nDrei\nVier")
    select_all(e)
    e.apply_style("Überschrift 2")                           # mehrere Absätze auf einmal
    assert all(b.blockFormat().headingLevel() == 2 for b in _blk(e.document()))
    c = e.textCursor(); c.setPosition(0); e.setTextCursor(c); e.apply_style("Standard")
    try:
        stile.aendern("Überschrift 2", groesse=30, farbe="#336699", fett=False, nach=20)
        n = e.stil_neu_anwenden("Überschrift 2")
        assert n == 3
        b = e.document().findBlockByNumber(1)
        cf = QTextCursor(b).charFormat()
        c2 = QTextCursor(b); c2.movePosition(QTextCursor.EndOfBlock)
        assert c2.charFormat().fontPointSize() == 30 and c2.charFormat().foreground().color().name() == "#336699"
        assert abs(b.blockFormat().bottomMargin() - 20 / 0.75) < 0.01
        assert e.document().findBlockByNumber(0).blockFormat().headingLevel() == 0          # „Standard“ unberührt
        assert "Überschrift 2" in stile.als_json()
        stile.laden(stile.als_json())
        assert stile.definition("Überschrift 2")["groesse"] == 30
    finally:
        stile.zuruecksetzen()
    assert stile.definition("Überschrift 2")["groesse"] == 16
    # „Kein Abstand“ ist eine eigene Vorlage und überlebt DOCX („No Spacing“)
    e2 = make(app, "Text"); e2.apply_style("Kein Abstand")
    assert e2.current_style() == "Kein Abstand"


def _blk(doc):
    b = doc.firstBlock()
    while b.isValid():
        yield b
        b = b.next()
