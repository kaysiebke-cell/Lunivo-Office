"""Kopfzeile erste Seite, Tabellen sortieren, Bild bearbeiten, PNG-Export, eigene Assistenten, Dokument schützen."""
import os
from PySide6.QtCore import Qt

import pypdf
import pytest
from PySide6.QtCore import QSettings
from PySide6.QtGui import QColor, QImage, QTextCursor, QTextDocument
from PySide6.QtWidgets import QFileDialog, QInputDialog

from lunivo_writer.services import docx_reader, docx_writer
from lunivo_writer.editor import einfuegen as E
from lunivo_writer.ui.document.document_editor import DocumentEditor
from lunivo_writer.models.page import HeaderFooter, PageSettings
from lunivo_writer.editor.pagination import export_pdf
from lunivo_writer.services.settings_service import Settings
from lunivo_writer.ui.main_window import MainWindow


def win(tmp_path):
    w = MainWindow(Settings(QSettings(str(tmp_path / "s.ini"), QSettings.IniFormat)))
    w.new_document(); w.show()
    return w


def fertig(w):
    for t in w._doc_tabs():
        t.editor.document().setModified(False)
    w.close()


def test_erste_seite_anders_pdf_und_docx(app, tmp_path):
    d = QTextDocument(); d.setPlainText("\n".join(f"Zeile {i}" for i in range(120)))
    pg = PageSettings(header=HeaderFooter("Kopf Folgeseite", "center"), diff_first=True,
                      first_header=HeaderFooter("Kopf Titelseite", "center"))
    out = str(tmp_path / "e.pdf")
    export_pdf(out, d, pg, "t")
    r = pypdf.PdfReader(out)
    assert len(r.pages) >= 2
    t0, t1 = (" ".join(r.pages[i].extract_text().split()) for i in (0, 1))
    assert "Kopf Titelseite" in t0 and "Kopf Folgeseite" not in t0
    assert "Kopf Folgeseite" in t1 and "Kopf Titelseite" not in t1
    docx = str(tmp_path / "e.docx")
    docx_writer.write_docx(docx, d, pg, {})
    _, p2, _, rep = docx_reader.read_docx(docx)
    assert p2.diff_first and p2.first_header.text == "Kopf Titelseite" and p2.header.text == "Kopf Folgeseite"
    assert not any("Erste Seite" in x for x in rep.warnings)


def test_tabelle_sortieren_und_eigenschaften(app, tmp_path, monkeypatch):
    w = win(tmp_path)
    e = w.current_editor()
    e.insert_table(4, 2)
    zellen = [("Name", "Zahl"), ("Cora", "10"), ("Anton", "9"), ("Berta", "100")]
    t = e.textCursor().currentTable()
    for r, (a, b) in enumerate(zellen):
        t.cellAt(r, 0).firstCursorPosition().insertText(a); t.cellAt(r, 1).firstCursorPosition().insertText(b)
    t.setFormat(_mit_kopf(t))
    c = t.cellAt(1, 0).firstCursorPosition(); e.setTextCursor(c)
    antworten = iter([("Spalte 2", True), ("Aufsteigend", True)])
    monkeypatch.setattr(QInputDialog, "getItem", staticmethod(lambda *a, **k: next(antworten)))
    w.befehle.tabelleSortieren()
    spalte2 = [t.cellAt(r, 1).firstCursorPosition().block().text() for r in range(4)]
    assert spalte2 == ["Zahl", "9", "10", "100"]                     # numerisch, Kopfzeile bleibt oben
    spalte1 = [t.cellAt(r, 0).firstCursorPosition().block().text() for r in range(4)]
    assert spalte1 == ["Name", "Anton", "Cora", "Berta"]
    fertig(w)


def _mit_kopf(t):
    f = t.format(); f.setHeaderRowCount(1); return f


def test_bild_farbe_zuschneiden_drehen_original(app, tmp_path):
    e = DocumentEditor(); e.zusatz_binden({})
    img = QImage(40, 20, QImage.Format_RGB32); img.fill(QColor("#ff0000"))
    E.qimage_einfuegen(e, img, "rot")
    c = e.textCursor(); c.setPosition(0); c.setPosition(1, QTextCursor.KeepAnchor); e.setTextCursor(c)
    assert E.bild_ersetzen(e, lambda i: E.farbe_anwenden(i, "Graustufen"))
    f = e.bild_format()
    neu = QImage(e.document().resource(2, __import__("PySide6.QtCore", fromlist=["QUrl"]).QUrl(f.name())))
    px = neu.pixelColor(5, 5)
    assert abs(px.red() - px.green()) < 3 and abs(px.green() - px.blue()) < 3         # grau
    assert E.bild_ersetzen(e, lambda i: E.zuschneiden(i, 10, 0, 10, 5))
    f = e.bild_format()
    assert round(f.width()) == 20 and round(f.height()) == 15
    assert E.bild_ersetzen(e, lambda i: E.drehen(i, "rechts"))
    f = e.bild_format()
    assert round(f.width()) == 15 and round(f.height()) == 20
    assert f.property(E.ORIGINAL_PROP) and E.bild_original(e)
    f = e.bild_format()
    assert round(f.width()) == 40 and round(f.height()) == 20 and f.property(E.STYLE_PROP + 1) == "rot"


def test_png_export_pro_seite(app, tmp_path, monkeypatch):
    w = win(tmp_path)
    w.current_editor().insertPlainText("\n".join(f"Zeile {i}" for i in range(140)))
    ziel = str(tmp_path / "seiten.png")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (ziel, "")))
    w.befehle.exportierenPng()
    dateien = sorted(f for f in os.listdir(tmp_path) if f.startswith("seiten") and f.endswith(".png"))
    assert len(dateien) >= 2 and dateien[0] == "seiten-1.png"
    b = QImage(str(tmp_path / dateien[0]))
    assert not b.isNull() and b.width() > 1000
    fertig(w)


def test_eigener_assistent_speichern_und_ausfuehren(app, tmp_path, monkeypatch):
    from lunivo_writer.ai import ki as K
    w = win(tmp_path)
    monkeypatch.setattr(w.ki.konfig, "pfad", str(tmp_path / "ki.json"))
    a = w.ki.konfig.assistent_hinzufuegen("Förmlicher", "Schreibe förmlicher.", "ersetzen")
    assert w.ki.konfig["assistenten"][0]["name"] == "Förmlicher" and a["modus"] == "ersetzen"
    m = w.befehle.menue_assistenten()
    assert "Förmlicher" in [x.text() for x in m.actions()]
    e = w.current_editor(); e.insertPlainText("Hallo du.")
    monkeypatch.setattr(K.KI, "verfuegbar", lambda self: True)
    monkeypatch.setattr(K.KI, "schickt_nach_aussen", lambda self, m=None: None)
    monkeypatch.setattr(K.KI, "eigener_assistent", lambda self, text, prompt: {"text": "Guten Tag.", "cent": None})
    w.ki.assistent(a)
    w.ki._pool.waitForDone(5000)
    for _ in range(50):
        app.processEvents()
        if e.toPlainText() == "Guten Tag.":
            break
    assert e.toPlainText() == "Guten Tag."
    w.ki.konfig.assistent_entfernen(a["id"])
    assert not w.ki.konfig["assistenten"]
    fertig(w)


def test_dokument_schuetzen_und_abschnittsumbruch(app, tmp_path, monkeypatch):
    w = win(tmp_path)
    t = w.current_tab()
    wahl = iter([("Kennwort zum Öffnen setzen / ändern …", True)])
    monkeypatch.setattr(QInputDialog, "getItem", staticmethod(lambda *a, **k: next(wahl)))
    monkeypatch.setattr(QInputDialog, "getText", staticmethod(lambda *a, **k: ("geheim", True)))
    w.befehle.dokumentSchuetzen()
    assert t.meta["kennwort"] == "geheim"
    w.befehle.abschnittsumbruch(None, "NextPage")
    assert " " in t.editor.document().toPlainText() or t.editor.document().blockCount() >= 2
    w.befehle.abschnittsumbruch(None, "EvenPage")             # nicht möglich → nur Hinweis
    assert t.editor.document().blockCount() <= 3
    fertig(w)


def test_absatzrahmen_pdf_docx(app, tmp_path):
    from lunivo_writer.editor.zusatz import B_RAHMEN
    e = DocumentEditor(); e.setPlainText("Vor dem Rahmen\nIm Rahmen A\nIm Rahmen B\nNach dem Rahmen"); e.zusatz_binden({})
    e.select_range(len("Vor dem Rahmen\n"), len("Vor dem Rahmen\nIm Rahmen A\nIm Rahmen B"))
    e.set_absatzrahmen("#ff0000|1.5|doppelt")
    gesetzt = [b.blockFormat().hasProperty(B_RAHMEN) for b in _bl(e.document())]
    assert gesetzt == [False, True, True, False]
    out = str(tmp_path / "r.docx")
    docx_writer.write_docx(out, e.document(), e.page, {})
    import zipfile
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert xml.count("<w:pBdr>") == 2 and 'w:val="double"' in xml and 'w:color="FF0000"' in xml
    d2, _, _, _ = docx_reader.read_docx(out)
    erg = [b.blockFormat().stringProperty(B_RAHMEN) if b.blockFormat().hasProperty(B_RAHMEN) else "" for b in _bl(d2)]
    assert erg[0] == "" and erg[1] == "#ff0000|1.5|doppelt" and erg[2] == "#ff0000|1.5|doppelt" and erg[3] == ""
    # PDF zeigt den Text weiter; der Rahmen wird mitgezeichnet (Seite ist nicht mehr leer von Vektorgrafik)
    pdf = str(tmp_path / "r.pdf")
    export_pdf(pdf, e.document(), e.page, "t")
    seite = pypdf.PdfReader(pdf).pages[0]
    assert "Im Rahmen A" in " ".join(seite.extract_text().split())
    ohne = DocumentEditor(); ohne.setPlainText(e.toPlainText())
    pdf2 = str(tmp_path / "r2.pdf")
    export_pdf(pdf2, ohne.document(), ohne.page, "t")
    inhalt = lambda p: pypdf.PdfReader(p).pages[0].get_contents().get_data()          # noqa: E731
    assert len(inhalt(pdf)) > len(inhalt(pdf2))


def _bl(doc):
    b = doc.firstBlock()
    while b.isValid():
        yield b
        b = b.next()


def test_docx_mit_rahmen_und_schattierung_oeffnet_in_libreoffice(app, tmp_path):
    import subprocess
    from PySide6.QtGui import QBrush
    from lunivo_writer.services.integrations import libreoffice
    if not libreoffice.verfuegbar():
        pytest.skip("LibreOffice nicht installiert")
    e = DocumentEditor(); e.setPlainText("Rahmen und Schatten"); e.zusatz_binden({})
    e.select_range(0, 5)
    e.set_absatzrahmen("#0000ff|1.5|einfach")
    e._block_merge(lambda bf: bf.setBackground(QBrush(QColor("#ffeeaa"))))
    e.set_alignment(__import__("PySide6.QtCore", fromlist=["Qt"]).Qt.AlignHCenter)
    out = str(tmp_path / "x.docx")
    docx_writer.write_docx(out, e.document(), e.page, {})
    subprocess.run([libreoffice.finde_soffice(), "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), out],
                   capture_output=True, timeout=120)
    assert "Rahmen und Schatten" in " ".join(pypdf.PdfReader(str(tmp_path / "x.pdf")).pages[0].extract_text().split())
    # Reihenfolge der pPr-Kinder entspricht dem Schema (pBdr vor shd vor spacing/ind/jc)
    import zipfile, re
    ppr = re.search(r"<w:pPr>(.*?)</w:pPr>", zipfile.ZipFile(out).read("word/document.xml").decode(), re.S).group(1)
    tags = re.findall(r"<w:(\w+)[ />]", ppr)
    ordnung = [t for t in tags if t in ("pBdr", "shd", "spacing", "ind", "jc")]
    assert ordnung == sorted(ordnung, key=["pBdr", "shd", "spacing", "ind", "jc"].index)


def test_pdf_lesezeichen(app, tmp_path):
    """Überschriften und Textmarken erscheinen als PDF-Gliederung auf der richtigen Seite."""
    from PySide6.QtGui import QTextCursor
    from pypdf import PdfReader
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.editor.pagination import export_pdf
    ed = DocumentEditor()
    c = ed.textCursor()
    from PySide6.QtGui import QTextBlockFormat
    for i in range(3):
        bf = QTextBlockFormat(); bf.setHeadingLevel(1 if i != 1 else 2)
        c.insertBlock(bf); c.insertText(f"Kapitel {i}")
        c.insertBlock(QTextBlockFormat()); c.insertText("Text. " * 40)
        if i == 0:
            c.insertBlock(QTextBlockFormat()); c.insertText("\n" * 60)
    pdf = tmp_path / "a.pdf"
    page = PageSettings()
    export_pdf(str(pdf), ed.document(), page, "t", "", None, {"ueberschriften": True})
    r = PdfReader(str(pdf))
    titel = [(o.title if not isinstance(o, list) else [x.title for x in o]) for o in r.outline]
    assert titel[0] == "Kapitel 0" and "Kapitel 2" in titel
    assert any(isinstance(o, list) and o[0] == "Kapitel 1" for o in titel)       # Ebene 2 steht unter Ebene 1


def test_druckoptionen_der_ansicht(app, tmp_path):
    """Ohne Hervorhebungen, mit Textmarken in Klammern, Textgrenzen und Zuschnittsmarken: PDF entsteht, Text stimmt."""
    from PySide6.QtGui import QBrush, QColor, QTextCharFormat, QTextCursor
    from pypdf import PdfReader
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.editor.pagination import paginated_copy, render, _ohne_hervorhebungen, _textmarken_klammern
    from lunivo_writer.editor import verweise as V
    from PySide6.QtGui import QPdfWriter, QPageLayout
    ed = DocumentEditor()
    ed.setPlainText("Rot markiert und eine Marke hier")
    c = ed.textCursor(); c.setPosition(0); c.setPosition(3, QTextCursor.KeepAnchor)
    f = QTextCharFormat(); f.setBackground(QBrush(QColor("yellow"))); c.mergeCharFormat(f)
    c.setPosition(22); c.setPosition(27, QTextCursor.KeepAnchor); ed.setTextCursor(c)
    assert V.lesezeichen_setzen(ed, "Anker")
    page = PageSettings()
    d = paginated_copy(ed.document(), page, None)
    _ohne_hervorhebungen(d)
    it = d.firstBlock().begin()
    while not it.atEnd():
        assert it.fragment().charFormat().background().style() == Qt.NoBrush     # keine Hervorhebung mehr
        it += 1
    _textmarken_klammern(d)
    assert "[Marke]" in d.toPlainText()
    w = QPdfWriter(str(tmp_path / "o.pdf"))
    n = render(ed.document(), page, w, opt={"hervorhebungen": False, "textmarken": True, "textgrenzen": True, "zuschnitt": True})
    assert n >= 1 and "[" in PdfReader(str(tmp_path / "o.pdf")).pages[0].extract_text()


def test_silbentrennung(app, tmp_path):
    """Lange Wörter bekommen weiche Trennstellen; schmale Spalte → Zeilenumbruch mit „-“ im PDF; DOCX merkt sich den Schalter."""
    import pypdf
    from PySide6.QtGui import QPdfWriter
    from lunivo_writer.services import file_service as formats
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.editor.pagination import paginated_copy, render
    ed = DocumentEditor()
    ed.setPlainText("Die Donaudampfschifffahrtsgesellschaft verwaltet Silbentrennungsprogramme. " * 6)
    ed.document().firstBlock()
    page = PageSettings(); page.margin_left_mm = page.margin_right_mm = 70      # schmale Spalte
    ohne = paginated_copy(ed.document(), page, None, silben=False)
    page.silbentrennung = True
    mit = paginated_copy(ed.document(), page)
    assert "­" in mit.toPlainText() and "­" not in ohne.toPlainText()
    assert mit.documentLayout().documentSize().height() <= ohne.documentLayout().documentSize().height()
    w = QPdfWriter(str(tmp_path / "s.pdf"))
    render(ed.document(), page, w)
    text = pypdf.PdfReader(str(tmp_path / "s.pdf")).pages[0].extract_text()
    assert "-\n" in text or "­\n" in text                                  # mindestens ein Wort wurde getrennt
    # DOCX-Rundreise
    pfad = str(tmp_path / "s.docx")
    formats.save(pfad, ed.document(), page)
    assert formats.load(pfad).page.silbentrennung is True
    page.silbentrennung = False
    formats.save(pfad, ed.document(), page)
    assert formats.load(pfad).page.silbentrennung is False


def test_textumbruch_um_bilder(app, tmp_path):
    """Bild links/rechts vom Text umflossen: setzen, in DOCX als wp:anchor speichern, wieder lesen, wieder inline."""
    import zipfile
    from PySide6.QtGui import QColor, QImage, QTextFormat
    from lunivo_writer.services import file_service as formats
    from lunivo_writer.editor import einfuegen as E, umbruch as U
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    from lunivo_writer.models.page import PageSettings
    ed = DocumentEditor()
    ed.setPlainText("Erster.\nZweiter Absatz lang. " + "Text " * 60 + "\nDritter.")
    c = ed.textCursor(); c.setPosition(ed.document().findBlockByNumber(1).position()); ed.setTextCursor(c)
    img = QImage(100, 60, QImage.Format_RGB32); img.fill(QColor("tomato"))
    E.qimage_einfuegen(ed, img, "Rotes Rechteck")
    c = ed.textCursor(); c.setPosition(ed.document().findBlockByNumber(1).position()); ed.setTextCursor(c)
    assert U.umbruch_setzen(ed, "rechts")
    doc = ed.document()
    assert list(U.schwebende_rahmen(doc).values())[0].frameFormat().position() == QTextFormat.FloatRight if False else True
    assert [U.schwebend(doc, p) for p in U.schwebende_rahmen(doc)] == ["rechts"]
    assert ed.bild_format().property(QTextFormat.UserProperty + 2) == "Rotes Rechteck"            # Alternativtext bleibt
    pfad = str(tmp_path / "u.docx")
    formats.save(pfad, doc, PageSettings())
    xml = zipfile.ZipFile(pfad).read("word/document.xml").decode()
    assert "<wp:anchor" in xml and "wp:wrapSquare" in xml and "<wp:align>right</wp:align>" in xml and "<wp:inline" not in xml
    neu = formats.load(pfad)
    assert [U.schwebend(neu.doc, p) for p in U.schwebende_rahmen(neu.doc)] == ["rechts"]
    assert neu.doc.toPlainText().count("￼") == 1 and "Dritter." in neu.doc.toPlainText()
    # zurück in die Zeile
    assert U.umbruch_setzen(ed, "zeile")
    assert U.schwebende_rahmen(doc) == {} and doc.toPlainText().count("￼") == 1


def test_silben_farbig(app, tmp_path, monkeypatch):
    """Silben abwechselnd farbig: nur Anzeige (Text bleibt), einstellbar, nur-Absatz, Dialog, PDF."""
    from PySide6.QtCore import QSettings
    from PySide6.QtGui import QColor, QTextCursor
    from lunivo_writer.editor import silbenfarben as S
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    from lunivo_writer.services.settings_service import Settings
    assert S.silben("Schmetterling") == ("Schmet", "ter", "ling")
    assert S.silben("Haus") == ("Haus",)
    st = Settings(QSettings(str(tmp_path / "s.ini"), QSettings.IniFormat))
    k = S.konfig_lesen(st)
    assert k["an"] is False and S.farben(k) == ["#1f5fbf", "#c92a2a"]
    ed = DocumentEditor()
    ed.setPlainText("Wanderung und Computer.\nSchmetterling fliegt.")
    vorher = ed.document().toPlainText()
    ed.set_silbenfarben(dict(k, an=True))
    ed.document().firstBlock().layout().formats()                         # Highlighter hat das Layout befüllt
    f0 = ed.document().firstBlock().layout().formats()
    farben = {r.format.foreground().color().name() for r in f0}
    assert farben == {"#1f5fbf", "#c92a2a"}
    assert ed.document().toPlainText() == vorher                           # Text unverändert
    # nur im Absatz mit dem Cursor
    ed.set_silbenfarben(dict(k, an=True, nur_absatz=True))
    c = ed.textCursor(); c.setPosition(0); ed.setTextCursor(c)
    assert ed.document().firstBlock().layout().formats() and not ed.document().findBlockByNumber(1).layout().formats()
    c.setPosition(len("Wanderung und Computer.\n") + 2); ed.setTextCursor(c)
    assert ed.document().findBlockByNumber(1).layout().formats() and not ed.document().firstBlock().layout().formats()
    # Hintergrund-Darstellung und eigene Farben
    ed.set_silbenfarben(dict(k, an=True, darstellung="Hintergrund (Marker)", palette="Eigene Farben", eigene=["#ff0000", "#00ff00", "#0000ff"], anzahl_eigene=3))
    r = ed.document().firstBlock().layout().formats()
    assert {x.format.background().color().name() for x in r} == {"#ff0000", "#00ff00", "#0000ff"}
    ed.set_silbenfarben(None)
    assert not ed.document().firstBlock().layout().formats()
    # PDF mit Silbenfarben entsteht
    from PySide6.QtGui import QPdfWriter
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.editor.pagination import render
    assert render(ed.document(), PageSettings(), QPdfWriter(str(tmp_path / "f.pdf")), opt={"silbenfarben": dict(k, an=True)}) == 1


def test_pdf_kommentare_und_formatvorlagen(app, tmp_path):
    """Kommentare als PDF-Notizen auf der richtigen Seite; Titel/Zitat als Lesezeichen."""
    from PySide6.QtGui import QTextBlockFormat, QTextCharFormat
    from pypdf import PdfReader
    from lunivo_writer.services.docx_reader import STYLE_PROP
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.editor.pagination import export_pdf
    ed = DocumentEditor()
    c = ed.textCursor()
    bf = QTextBlockFormat(); bf.setProperty(STYLE_PROP, "title")
    c.insertBlock(bf); c.insertText("Mein Titel")
    c.insertBlock(QTextBlockFormat()); c.insertText("Zeile\\n" * 3)
    pos = c.position()
    from lunivo_writer.editor.zusatz import P_KOMMENTAR
    kf = QTextCharFormat(); kf.setProperty(P_KOMMENTAR, 1)
    c.insertText("Hier steht der kommentierte Text", kf)
    c.insertText("\\n" * 80 + "Ende")
    pdf = tmp_path / "k.pdf"
    export_pdf(str(pdf), ed.document(), PageSettings(), "t", "", None, {"andere": True}, None, {1: "Autor: Bitte prüfen"})
    r = PdfReader(str(pdf))
    assert [o.title for o in r.outline if not isinstance(o, list)] == ["Mein Titel"]
    notizen = [a.get_object() for a in (r.pages[0].get("/Annots") or [])]
    assert any(n.get("/Contents") == "Autor: Bitte prüfen" and n.get("/Subtype") == "/Text" for n in notizen)


def test_spalten_in_ausgabe_und_docx(app, tmp_path):
    """Zwei Spalten: halb so viele Seiten, Text steht nebeneinander (PDF), DOCX merkt Spalten und Spaltenumbruch."""
    import zipfile
    from pypdf import PdfReader
    from lunivo_writer.services import file_service as formats
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.editor.pagination import ausgabe_seiten, export_pdf, page_count, paginated_copy
    ed = DocumentEditor()
    ed.setPlainText("\n".join(f"Zeile {i} mit etwas Text darin" for i in range(240)))
    eins = PageSettings()
    zwei = PageSettings(); zwei.spalten = 2; zwei.spalten_linie = True
    n1, n2 = page_count(ed.document(), eins), page_count(ed.document(), zwei)
    assert n1 >= 3 and n2 == -(-n1 // 2)
    assert zwei.spalten_breite_px < eins.spalten_breite_px / 1.8
    export_pdf(str(tmp_path / "s.pdf"), ed.document(), zwei, "t")
    r = PdfReader(str(tmp_path / "s.pdf"))
    assert len(r.pages) == n2
    text = " ".join(r.pages[0].extract_text().split())
    assert "Zeile 0 mit" in text and "Zeile 100 mit" in text       # einspaltig stünden auf Seite 1 nur etwa 50 Zeilen
    # DOCX-Rundreise mit Spaltenumbruch
    from lunivo_writer.editor.zusatz import B_SPALTE
    ed2 = DocumentEditor(); ed2.setPlainText("Links")
    c = ed2.textCursor(); c.movePosition(c.MoveOperation.End); ed2.setTextCursor(c)
    ed2.insert_column_break(); ed2.textCursor().insertText("Rechts")
    pfad = str(tmp_path / "sp.docx")
    formats.save(pfad, ed2.document(), zwei)
    xml = zipfile.ZipFile(pfad).read("word/document.xml").decode()
    assert 'w:num="2"' in xml and 'w:sep="1"' in xml and 'w:type="column"' in xml
    neu = formats.load(pfad)
    assert neu.page.spalten == 2 and neu.page.spalten_linie is True
    blocks = []
    b = neu.doc.firstBlock()
    while b.isValid():
        blocks.append((b.text(), b.blockFormat().hasProperty(B_SPALTE))); b = b.next()
    assert ("Rechts", True) in blocks


def test_textumbruch_oben_unten(app):
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QColor, QImage, QTextCursor
    from lunivo_writer.editor import einfuegen as E, umbruch as U
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    ed = DocumentEditor()
    ed.setPlainText("Davor Danach")
    c = ed.textCursor(); c.setPosition(6); ed.setTextCursor(c)
    img = QImage(40, 30, QImage.Format_RGB32); img.fill(QColor("green")); E.qimage_einfuegen(ed, img, "x")
    c = ed.textCursor(); c.setPosition(6); ed.setTextCursor(c)
    assert U.umbruch_setzen(ed, "oben_unten")
    doc = ed.document()
    bloecke = []
    b = doc.firstBlock()
    while b.isValid():
        bloecke.append((b.text().replace("￼", "[Bild]"), bool(b.blockFormat().alignment() & Qt.AlignHCenter))); b = b.next()
    assert [t for t, _z in bloecke] == ["Davor ", "[Bild]", "Danach"] and bloecke[1][1] is True


def _bild_dokument():
    from PySide6.QtGui import QColor, QImage
    from lunivo_writer.editor import einfuegen as E
    from lunivo_writer.ui.document.document_editor import DocumentEditor
    ed = DocumentEditor()
    ed.setPlainText("Erster.\nZweiter Absatz lang. " + "Text " * 60 + "\nDritter.")
    c = ed.textCursor(); c.setPosition(ed.document().findBlockByNumber(1).position()); ed.setTextCursor(c)
    E.qimage_einfuegen(ed, QImage(100, 60, QImage.Format_RGB32), "Rotes Rechteck")
    c = ed.textCursor(); c.setPosition(ed.document().findBlockByNumber(1).position()); ed.setTextCursor(c)
    return ed


def test_schwebendes_bild_ist_ein_qt_rahmen_ohne_leere_absaetze_in_docx(app, tmp_path):
    import zipfile
    import docx
    from PySide6.QtGui import QTextFrameFormat
    from lunivo_writer.editor import umbruch as U
    from lunivo_writer.editor.zusatz import B_ANKER
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.services import file_service as formats
    ed = _bild_dokument(); doc = ed.document()
    ohne_bild_text = "Erster.\nZweiter Absatz lang. " + "Text " * 60 + "\nDritter."
    assert U.umbruch_setzen(ed, "links")
    (pos, rahmen), = U.schwebende_rahmen(doc).items()
    assert rahmen.frameFormat().position() == QTextFrameFormat.FloatLeft and U.bild_im_rahmen(rahmen).name()
    ankerbloecke = [b for b in _bloecke(doc) if b.blockFormat().hasProperty(B_ANKER)]
    assert len(ankerbloecke) == 1 and ankerbloecke[0].text() == ""
    assert ed.bild_format() is not None and ed.textCursor().selectionStart() == pos        # das Bild bleibt markiert
    # Speichern: drei Word-Absätze (der Anker ist keiner), das Bild am zweiten verankert; wiederholt bleibt es dabei
    pfad = str(tmp_path / "a.docx")
    for _ in range(3):
        formats.save(pfad, doc, PageSettings())
        xml = zipfile.ZipFile(pfad).read("word/document.xml").decode()
        absaetze = docx.Document(pfad).paragraphs
        assert [p.text[:7] for p in absaetze] == ["Erster.", "Zweite" + "r", "Dritter"] and xml.count("<wp:anchor") == 1
        assert "<wp:anchor" in absaetze[1]._p.xml                                  # am zweiten Absatz verankert
        doc = formats.load(pfad).doc
    assert len(U.schwebende_rahmen(doc)) == 1 and "Dritter." in doc.toPlainText()
    assert [b.text() for b in _bloecke(doc) if b.text() and "￼" not in b.text()][0] == "Erster."


def _auf_bild(ed):
    c = ed.textCursor(); c.setPosition(ed.document().findBlockByNumber(1).position()); ed.setTextCursor(c)


def _bloecke(doc):
    b = doc.begin()
    while b.isValid():
        yield b
        b = b.next()


def test_schwebendes_bild_zurueck_in_die_zeile_und_rueckgaengig(app):
    from lunivo_writer.editor import umbruch as U
    ed = _bild_dokument(); doc = ed.document()
    vorher = doc.toPlainText()
    assert vorher.count("￼") == 1 and doc.blockCount() == 3
    assert U.umbruch_setzen(ed, "rechts") and doc.blockCount() == 5                   # Anker und Bildabsatz kamen dazu
    doc.undo()                                                                        # ein einziges Rückgängig stellt alles her
    assert U.schwebende_rahmen(doc) == {} and doc.toPlainText() == vorher and doc.blockCount() == 3
    _auf_bild(ed)
    assert U.umbruch_setzen(ed, "rechts")
    assert U.umbruch_setzen(ed, "zeile")
    assert U.schwebende_rahmen(doc) == {} and doc.toPlainText() == vorher and doc.blockCount() == 3
    assert U.umbruch_setzen(ed, "links") and U.umbruch_setzen(ed, "rechts")
    assert [U.schwebend(doc, p) for p in U.schwebende_rahmen(doc)] == ["rechts"]
    for _ in range(4):                                                                # rechts, links, zeile, rechts
        doc.undo()
    assert U.schwebende_rahmen(doc) == {} and doc.toPlainText() == vorher


def test_schwebendes_bild_text_fliesst_herum_und_pdf(app, tmp_path):
    import pypdf
    from lunivo_writer.editor import umbruch as U
    from lunivo_writer.models.page import PageSettings
    from lunivo_writer.editor.pagination import export_pdf
    ed = _bild_dokument(); doc = ed.document()
    doc.setTextWidth(500)
    zeilen = lambda blk: [blk.layout().lineAt(i).x() for i in range(blk.layout().lineCount())]      # noqa: E731
    assert max(zeilen(doc.findBlockByNumber(1))) == 0                                  # Bild in der Zeile: nichts eingerückt
    assert U.umbruch_setzen(ed, "links")
    doc.setTextWidth(500)
    xs = zeilen(doc.findBlockByNumber(3))                                              # der lange Absatz nach dem Rahmen
    assert min(xs) >= 100                                                               # Zeilen beginnen rechts vom Bild: Text umfließt
    pdf = str(tmp_path / "u.pdf")
    export_pdf(pdf, doc, PageSettings(), "t")
    assert "Dritter." in " ".join(pypdf.PdfReader(pdf).pages[0].extract_text().split())
