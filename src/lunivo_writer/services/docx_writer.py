"""QTextDocument → DOCX (python-docx).

Schreibt Absätze, Überschriften (Formatvorlagen „Heading n“), Titel/Untertitel/
Zitat, Zeichenformat, Ausrichtung, Einzüge, Abstände, Zeilenabstand, Listen,
Seitenumbrüche, Tabellen (mit verbundenen Zellen), Bilder, Hyperlinks,
Seitengröße, Ränder, Kopf- und Fußzeile (mit Seitenzahlfeld).
"""
from __future__ import annotations

import io
import os
import tempfile
from typing import Optional

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor
from PySide6.QtCore import QBuffer, QIODevice, QUrl, Qt
from PySide6.QtGui import (
    QFont, QImage, QTextBlock, QTextBlockFormat, QTextCharFormat, QTextDocument, QTextFormat,
    QTextFrame, QTextListFormat, QTextTable,
)

from lunivo_writer.models.page import PageSettings
from lunivo_writer.services.docx_reader import LH_FIXED, LH_PROPORTIONAL, STYLE_PROP
from lunivo_writer.models.report import DocumentIOError, IOReport

PX_TO_PT = 0.75
EMU_PER_PX = 9525

_LIST_STYLES = {
    QTextListFormat.ListDisc: "bullet", QTextListFormat.ListCircle: "bullet",
    QTextListFormat.ListSquare: "bullet", QTextListFormat.ListDecimal: "number",
    QTextListFormat.ListLowerAlpha: "number", QTextListFormat.ListUpperAlpha: "number",
    QTextListFormat.ListLowerRoman: "number", QTextListFormat.ListUpperRoman: "number",
}


def _has(fmt, prop) -> bool:
    return fmt.hasProperty(prop)


from lunivo_writer.editor.zusatz import B_SPALTE, P_OBJEKT, P_VERSTECKT  # noqa: E402
from lunivo_writer.editor.formular import P_FORM  # noqa: E402
from lunivo_writer.editor.abschnitte import B_ABSCHNITT  # noqa: E402
from lunivo_writer.editor import umbruch
from lunivo_writer.editor.linie import Linie, art_aus_docx, docx_wert
from lunivo_writer.editor import muster
from lunivo_writer.editor.zusatz import B_ART, B_CAP, B_MUSTER, B_RAHMEN, B_TOC, P_AENDERUNG, P_AUTOR, P_DATUM, P_FELD, P_INDEX, P_KOMMENTAR, P_NOTE


class _Writer:
    def __init__(self, qdoc: QTextDocument, page: PageSettings, meta: dict):
        self.q = qdoc
        self.page = page
        self.meta = meta
        self.report = IOReport()
        self.d = docx.Document()
        self._clear_body()
        self._num_ids: dict[tuple, int] = {}
        self._vorab_bilder: list = []                    # schwebende Bilder, die am nächsten Absatz verankert werden
        self.last_list_obj = None
        self.last_num = None
        self._komm_runs: dict[int, list] = {}
        self._noten: dict[str, list] = {"f": [], "e": []}       # [(Ausgabe-Id, Text)]
        self._note_ids: dict[tuple, int] = {}
        self._aend_id = 0
        self._lz_id = 0
        self._lz_offen: dict[str, int] = {}

    def _clear_body(self):
        body = self.d.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)

    # ---- Seite --------------------------------------------------------
    def write_page(self):
        """Dokumentweite Einstellungen (Hintergrund, Silbentrennung) und der letzte Abschnitt in der Schlusseinrichtung des Textes."""
        from lunivo_writer.editor import abschnitte as AB
        self.abschnitte = AB.alle_einrichtungen(self.q, self.page)
        self._ab_nr = 0
        self._dokument_ebene(self.page)
        typ, letzte = self.abschnitte[-1]
        self.write_section(self.d.sections[0], letzte, typ if len(self.abschnitte) > 1 else None, len(self.abschnitte) - 1)

    def write_section(self, s, p, typ: str | None = None, nr: int = 0):
        """Einen Abschnitt (python-docx Section) nach Einrichtung p schreiben; nr = Nummer des Abschnitts (für verknüpfte Kopfzeilen)."""
        from lunivo_writer.editor import abschnitte as AB
        s.page_width, s.page_height = Mm(p.width_mm), Mm(p.height_mm)
        s.orientation = WD_ORIENT.LANDSCAPE if p.landscape else WD_ORIENT.PORTRAIT
        s.top_margin, s.bottom_margin = Mm(p.margin_top_mm), Mm(p.margin_bottom_mm)
        s.left_margin, s.right_margin = Mm(p.margin_left_mm), Mm(p.margin_right_mm)
        if typ:
            from docx.enum.section import WD_SECTION
            s.start_type = {"naechste": WD_SECTION.NEW_PAGE, "fortlaufend": WD_SECTION.CONTINUOUS,
                            "gerade": WD_SECTION.EVEN_PAGE, "ungerade": WD_SECTION.ODD_PAGE}.get(typ, WD_SECTION.NEW_PAGE)
        self._seitenschmuck(s, p)
        paare = [(s.header, p.header, p.kopf_verknuepft), (s.footer, p.footer, p.fuss_verknuepft)]
        if p.diff_first:
            s.different_first_page_header_footer = True
            paare += [(s.first_page_header, p.first_header, p.kopf_verknuepft), (s.first_page_footer, p.first_footer, p.fuss_verknuepft)]
        for part, hf, verknuepft in paare:
            if nr > 0 and verknuepft:
                continue                                            # bleibt „mit vorherigem verknüpft“
            if hf.is_empty() and nr == 0:
                continue
            part.is_linked_to_previous = False
            lines = hf.text.split("\n")
            for i, line in enumerate(lines):
                para = part.paragraphs[0] if i == 0 else part.add_paragraph()
                para.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                                  "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(hf.align, WD_ALIGN_PARAGRAPH.LEFT)
                self._text_with_fields(para, line)

    def _abschnitt_ende(self, para) -> None:
        """Am Ende des Absatzes, der einen Abschnitt beschließt: dessen Einrichtung als w:sectPr in den Absatz schreiben."""
        from docx.section import Section
        k = self._ab_nr
        self._ab_nr += 1
        typ, ps = self.abschnitte[k]
        sp = OxmlElement("w:sectPr")
        ppr = para._p.get_or_add_pPr()
        nach = ppr.find(qn("w:pPrChange"))
        (nach.addprevious(sp) if nach is not None else ppr.append(sp))
        self.write_section(Section(sp, self.d.part), ps, typ if k > 0 else None, k)

    def _dokument_ebene(self, p) -> None:
        doc_el = self.d.element
        for alt in doc_el.findall(qn("w:background")):
            doc_el.remove(alt)
        if p.color:
            bg = OxmlElement("w:background")
            bg.set(qn("w:color"), p.color.lstrip("#").upper())
            doc_el.insert(0, bg)
            st = self.d.settings.element
            if st.find(qn("w:displayBackgroundShape")) is None:
                st.insert(0, OxmlElement("w:displayBackgroundShape"))
        st = self.d.settings.element
        for alt in st.findall(qn("w:autoHyphenation")):
            st.remove(alt)
        if p.silbentrennung:
            ah = OxmlElement("w:autoHyphenation")
            anker = st.find(qn("w:zoom"))
            (anker.addnext(ah) if anker is not None else st.insert(0, ah))

    def _seitenschmuck(self, sect, p) -> None:
        from docx.oxml import OxmlElement
        sp = sect._sectPr
        for alt in sp.findall(qn("w:pgBorders")):
            sp.remove(alt)
        if p.border:
            b = p.border
            pb = OxmlElement("w:pgBorders")
            pb.set(qn("w:offsetFrom"), "page")
            for seite, kurz in (("top", "t"), ("left", "l"), ("bottom", "b"), ("right", "r")):
                if kurz not in b.get("seiten", "tblr"):
                    continue
                e = OxmlElement(f"w:{seite}")
                e.set(qn("w:val"), docx_wert(b.get("art", "einfach")))
                e.set(qn("w:sz"), str(max(2, int(round(b.get("breite_pt", 1.0) * 8)))))
                e.set(qn("w:space"), str(int(round(b.get("abstand_mm", 8.0) / 25.4 * 72))))
                e.set(qn("w:color"), b.get("farbe", "#000000").lstrip("#").upper())
                pb.append(e)
            pg_mar = sp.find(qn("w:pgMar"))
            (pg_mar.addnext(pb) if pg_mar is not None else sp.append(pb))
        for alt in sp.findall(qn("w:cols")):
            sp.remove(alt)
        if p.spalten > 1:
            co = OxmlElement("w:cols")
            co.set(qn("w:num"), str(int(p.spalten)))
            co.set(qn("w:space"), str(int(round(p.spalten_abstand_mm / 25.4 * 1440))))
            if p.spalten_linie:
                co.set(qn("w:sep"), "1")
            anker = sp.find(qn("w:pgMar"))
            (anker.addnext(co) if anker is not None else sp.append(co))
        for alt in sp.findall(qn("w:lnNumType")):
            sp.remove(alt)
        if p.zeilennummern:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
            anker = sp.find(qn("w:pgBorders"))
            if anker is None:
                anker = sp.find(qn("w:pgMar"))
            (anker.addnext(ln) if anker is not None else sp.append(ln))
        if p.watermark and p.watermark.get("text") and not getattr(self, "_wasserzeichen_gemeldet", False):
            self._wasserzeichen_gemeldet = True
            self.report.warn("Das Wasserzeichen wird in der DOCX-Datei nicht gespeichert (nur Anzeige, PDF und Druck).")

    def _text_with_fields(self, para, text: str):
        import re
        for tok in re.split(r"(\{PAGE\}|\{PAGES\})", text):
            if tok in ("{PAGE}", "{PAGES}"):
                self._field(para, "PAGE" if tok == "{PAGE}" else "NUMPAGES")
            elif tok:
                para.add_run(tok)

    @staticmethod
    def _field(para, name: str):
        run = para.add_run()
        for kind, extra in (("begin", None), (None, f" {name} "), ("separate", None), (None, "1"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                run._r.append(el)
            elif extra.strip() == name:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = extra
                run._r.append(it)
            else:
                t = OxmlElement("w:t")
                t.text = extra
                run._r.append(t)

    # ---- Grundschrift --------------------------------------------------
    def _sprache(self) -> None:
        kennung = self.meta.get("sprache")
        if not kennung:
            return
        rpr = self.d.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is None:
            return
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), kennung)

    def write_defaults(self):
        f: QFont = self.q.defaultFont()
        normal = self.d.styles["Normal"]
        normal.font.name = f.family()
        if f.pointSizeF() > 0:
            normal.font.size = Pt(f.pointSizeF())
        rpr = normal.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), f.family())
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        pf = normal.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)

    # ---- Durchlauf -----------------------------------------------------
    def write_frame(self, frame: QTextFrame, container):
        it = frame.begin()
        first_in_cell = isinstance(container, _CellBox)
        while not it.atEnd():
            blk = it.currentBlock()
            if blk.isValid() and umbruch.ist_anker(blk):
                it += 1                                          # Hilfsabsatz vor einem schwebenden Bild: kein Word-Absatz
                continue
            if blk.isValid():
                para = container.first_paragraph() if first_in_cell else container.add_paragraph()
                first_in_cell = False
                self.write_block(blk, para)
                if isinstance(container, _Body) and blk.blockFormat().hasProperty(B_ABSCHNITT):
                    self._abschnitt_ende(para)
            else:
                fr = it.currentFrame()
                if isinstance(fr, QTextTable):
                    self.write_table(fr, container)
                    first_in_cell = False
                elif fr is not None and fr.frameFormat().position() in umbruch.SEITEN and umbruch.bild_im_rahmen(fr) is not None:
                    self._vorab_bilder.append((umbruch.bild_im_rahmen(fr), umbruch.SEITEN[fr.frameFormat().position()]))
                elif fr is not None:
                    self.report.warn("Ein Rahmen im Text wurde nur als Fließtext gespeichert.")
                    self.write_frame(fr, container)
            it += 1

    def _eigener_stil(self, name: str):
        """Selbst angelegte Formatvorlage als echte Word-Absatzvorlage (Schrift, Größe, Farbe, Abstände, Ausrichtung)."""
        from lunivo_writer.editor import stile
        if not stile.ist_eigen(name):
            return None
        if name in [x.name for x in self.d.styles]:
            return self.d.styles[name]
        d = stile.definition(name)
        st = self.d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = self.d.styles["Normal"]
        if d["groesse"]:
            st.font.size = Pt(d["groesse"])
        if d["fett"]:
            st.font.bold = True
        if d["kursiv"]:
            st.font.italic = True
        if d["schrift"]:
            st.font.name = d["schrift"]
        if d["farbe"]:
            st.font.color.rgb = RGBColor.from_string(d["farbe"].lstrip("#").upper())
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(d["vor"]), Pt(d["nach"])
        if d.get("zeilen"):
            st.paragraph_format.line_spacing = float(d["zeilen"])
        if d.get("einzug"):
            st.paragraph_format.left_indent = Mm(float(d["einzug"]))
        if d["ausr"]:
            st.paragraph_format.alignment = {"links": WD_ALIGN_PARAGRAPH.LEFT, "mitte": WD_ALIGN_PARAGRAPH.CENTER,
                                             "rechts": WD_ALIGN_PARAGRAPH.RIGHT, "block": WD_ALIGN_PARAGRAPH.JUSTIFY}[d["ausr"]]
        return st

    def write_block(self, blk: QTextBlock, para):
        bf = blk.blockFormat()
        style_name = bf.property(STYLE_PROP) if _has(bf, STYLE_PROP) else None
        lvl = bf.headingLevel()
        lst = blk.textList()
        if bf.hasProperty(B_TOC):
            kind = bf.property(B_TOC)
            art = bf.property(B_ART) if bf.hasProperty(B_ART) else "toc"
            name = ("TOC Heading" if kind == "kopf" else
                    {"toc": f"TOC {kind[-1]}", "abb": "Table of Figures", "tab": "Table of Figures",
                     "index": "Index 1", "lit": "Bibliography"}.get(art, "TOC 1"))
            para.style = self._stil(name)
        elif bf.hasProperty(B_CAP):
            para.style = self.d.styles["Caption"]
        elif lvl:
            para.style = self.d.styles[f"Heading {lvl}"]
        elif style_name in ("Title", "Subtitle", "Quote", "No Spacing"):
            para.style = self.d.styles[style_name]
        elif lst is not None:
            self._apply_list(para, lst)
        elif style_name and str(style_name).startswith("x:") and self._eigener_stil(str(style_name)[2:]) is not None:
            para.style = self._eigener_stil(str(style_name)[2:])
        pf = para.paragraph_format
        a = bf.alignment()
        if a & Qt.AlignHCenter:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif a & Qt.AlignRight:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif a & Qt.AlignJustify:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Einzüge: bei Listen steckt der Listeneinzug in der Vorlage
        if lst is None:
            if bf.leftMargin():
                pf.left_indent = Emu(int(bf.leftMargin() * EMU_PER_PX))
            if bf.textIndent():
                pf.first_line_indent = Emu(int(bf.textIndent() * EMU_PER_PX))
        if bf.rightMargin():
            pf.right_indent = Emu(int(bf.rightMargin() * EMU_PER_PX))
        if bf.layoutDirection() == Qt.RightToLeft:
            ppr = para._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")
            ppr.insert_element_before(bidi, "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                                      "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                                      "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
        pf.space_before = Pt(bf.topMargin() * PX_TO_PT)
        pf.space_after = Pt(bf.bottomMargin() * PX_TO_PT)
        lh, lht = bf.lineHeight(), bf.lineHeightType()
        if int(lht) == LH_PROPORTIONAL and lh and lh != 100:
            pf.line_spacing = lh / 100.0
        elif int(lht) == LH_FIXED and lh:
            pf.line_spacing = Pt(lh * PX_TO_PT)
        if bf.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore:
            if bf.hasProperty(B_SPALTE):                              # Spaltenumbruch: <w:br w:type="column"/> am Absatzanfang
                from docx.enum.text import WD_BREAK
                para.add_run().add_break(WD_BREAK.COLUMN)
            else:
                pf.page_break_before = True
        tabs = bf.tabPositions()
        if tabs:
            from docx.enum.text import WD_TAB_ALIGNMENT
            from PySide6.QtGui import QTextOption
            art = {QTextOption.RightTab: WD_TAB_ALIGNMENT.RIGHT, QTextOption.CenterTab: WD_TAB_ALIGNMENT.CENTER,
                   QTextOption.DelimiterTab: WD_TAB_ALIGNMENT.DECIMAL}
            for t in tabs:
                pf.tab_stops.add_tab_stop(Emu(int(t.position * EMU_PER_PX)), art.get(t.type, WD_TAB_ALIGNMENT.LEFT))
        nach_shd = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                    "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
        if bf.hasProperty(B_RAHMEN):
            li = Linie.aus_text(bf.stringProperty(B_RAHMEN))
            bdr = OxmlElement("w:pBdr")
            val = docx_wert(li.art)
            for seite, kurz in (("top", "t"), ("left", "l"), ("bottom", "b"), ("right", "r"), ("between", "h")):
                if kurz not in li.seiten:
                    continue
                e = OxmlElement(f"w:{seite}")
                e.set(qn("w:val"), val)
                e.set(qn("w:sz"), str(max(2, int(round(li.breite_pt * 8)))))
                e.set(qn("w:space"), "4")
                e.set(qn("w:color"), (li.farbe or "#000000").lstrip("#").upper())
                bdr.append(e)
            if len(bdr):
                para._p.get_or_add_pPr().insert_element_before(bdr, "w:shd", *nach_shd)
        bg = bf.background()
        mu = muster.aus_text(bf.stringProperty(B_MUSTER)) if bf.hasProperty(B_MUSTER) else None
        if mu:                                                   # Muster über Füllfarbe: w:val = Musterart, w:color = Musterfarbe
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), muster.ooxml_von(mu[0])); shd.set(qn("w:color"), mu[1][1:].upper()); shd.set(qn("w:fill"), mu[2][1:].upper() if mu[2] else "auto")
            para._p.get_or_add_pPr().insert_element_before(shd, *nach_shd)
        elif bg.style() != Qt.NoBrush and bg.color().isValid():
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), bg.color().name()[1:].upper())
            para._p.get_or_add_pPr().insert_element_before(shd, *nach_shd)
        # Läufe
        for imf, seite in self._vorab_bilder:                    # schwebende Bilder vor diesem Absatz: hier verankern
            self._image(imf, para, seite)
        self._vorab_bilder = []
        it = blk.begin()
        in_head = bool(lvl) or style_name in ("Title", "Subtitle")
        frags = []
        while not it.atEnd():
            if it.fragment().isValid():
                frags.append(it.fragment())
            it += 1
        i = 0
        while i < len(frags):
            cf = frags[i].charFormat()
            form = cf.stringProperty(P_FORM) if cf.hasProperty(P_FORM) else ""
            if form:                                                   # Formularfeld → Inhaltssteuerelement (w:sdt)
                j = i
                while j < len(frags) and frags[j].charFormat().hasProperty(P_FORM) and frags[j].charFormat().stringProperty(P_FORM) == form:
                    j += 1
                sdt, inhalt = self._steuerelement(form, "".join(f.text() for f in frags[i:j]))
                para._p.append(sdt)
                self._ziel = inhalt
                try:
                    for k in range(i, j):
                        self.write_fragment(frags[k], para, in_head)
                finally:
                    self._ziel = None
                i = j
                continue
            instr = cf.stringProperty(P_FELD) if cf.hasProperty(P_FELD) else ""
            if instr:
                j = i
                while j < len(frags) and frags[j].charFormat().hasProperty(P_FELD) and frags[j].charFormat().stringProperty(P_FELD) == instr:
                    j += 1
                self._textmarken(para, [])
                self._fldchar(para, "begin"); self._instr(para, instr); self._fldchar(para, "separate")
                for k in range(i, j):
                    self.write_fragment(frags[k], para, in_head)
                self._fldchar(para, "end")
                i = j
            else:
                self.write_fragment(frags[i], para, in_head)
                i += 1
        self._textmarken(para, [])

    def _stil(self, name: str):
        from docx.enum.style import WD_STYLE_TYPE
        if name not in [s.name for s in self.d.styles]:
            st = self.d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = self.d.styles["Normal"]
        return self.d.styles[name]

    @staticmethod
    def _fldchar(para, typ: str):
        r = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ)
        r.append(fc)
        para._p.append(r)

    @staticmethod
    def _instr(para, text: str):
        r = OxmlElement("w:r")
        it = OxmlElement("w:instrText"); it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); it.text = f" {text} "
        r.append(it)
        para._p.append(r)

    def _apply_list(self, para, lst):
        fmt = lst.format()
        kind = _LIST_STYLES.get(fmt.style(), "bullet")
        level = max(1, fmt.indent())
        base = "List Bullet" if kind == "bullet" else "List Number"
        name = base if level == 1 else f"{base} {min(level, 3)}"
        try:
            para.style = self.d.styles[name]
        except KeyError:
            para.style = self.d.styles[base]
        if kind == "number":
            key = lst.item(0).blockNumber()
            if key not in self._num_ids:
                self._num_ids[key] = self._new_num_instance(para)
            numpr = para._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = self._num_ids[key]

    def _new_num_instance(self, para) -> int:
        """Eigene Zählung je Liste, damit „1.“ neu beginnt."""
        numbering = self.d.part.numbering_part.element
        base = None
        for s in (para.style, para.style.base_style):
            if s is not None and s.element.pPr is not None and s.element.pPr.numPr is not None:
                base = s.element.pPr.numPr.numId.val
                break
        if base is None:
            return 1
        src = numbering.xpath(f'w:num[@w:numId="{base}"]')[0]
        abs_id = src.xpath("w:abstractNumId/@w:val")[0]
        new_id = max([int(n) for n in numbering.xpath("w:num/@w:numId")] + [0]) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), abs_id)
        num.append(an)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        numbering.append(num)
        return new_id

    # ---- Läufe ----------------------------------------------------------
    def write_fragment(self, frag, para, in_head: bool):
        cf: QTextCharFormat = frag.charFormat()
        text = frag.text()
        if cf.isImageFormat():
            self._image(cf.toImageFormat(), para)
            return
        parent = getattr(self, "_ziel", None) if getattr(self, "_ziel", None) is not None else para
        href = cf.anchorHref() if cf.isAnchor() else ""
        self._textmarken(para, [] if href else (list(cf.anchorNames()) if cf.isAnchor() else []))
        if href:
            hl = self._hyperlink(para, href)
            parent = hl
        note = cf.stringProperty(P_NOTE) if cf.hasProperty(P_NOTE) else ""
        if note:
            run = self._make_run(parent, para)
            self._note_verweis(run, note)
            self._zusatz_anwenden(run, cf)
            return
        for chunk, kind in _split(text):
            run = self._make_run(parent, para)
            if kind == "text":
                run.text = chunk
            elif kind == "tab":
                run.add_tab()
            elif kind == "br":
                run.add_break()
            elif kind == "img":
                continue
            self._format_run(run, cf, in_head, bool(href))
            self._zusatz_anwenden(run, cf)
        if cf.hasProperty(P_INDEX) and cf.stringProperty(P_INDEX):
            self._fldchar(para, "begin")
            self._instr(para, 'XE "' + cf.stringProperty(P_INDEX).replace('"', "'") + '"')
            self._fldchar(para, "end")

    def _textmarken(self, para, namen: list) -> None:
        """Offene Textmarken schließen, neue öffnen — je Name genau ein bookmarkStart/-End-Paar."""
        for n in [n for n in self._lz_offen if n not in namen]:
            e = OxmlElement("w:bookmarkEnd")
            e.set(qn("w:id"), str(self._lz_offen.pop(n)))
            para._p.append(e)
        for n in namen:
            if n not in self._lz_offen:
                self._lz_id += 1
                st = OxmlElement("w:bookmarkStart")
                st.set(qn("w:id"), str(self._lz_id)); st.set(qn("w:name"), n)
                para._p.append(st)
                self._lz_offen[n] = self._lz_id

    def _note_verweis(self, run, note: str) -> None:
        art, _, nid = note.partition(":")
        nid = int(nid)
        key = (art, nid)
        if key not in self._note_ids:
            self._note_ids[key] = len(self._noten[art]) + 1
            self._noten[art].append((self._note_ids[key], (self.meta.get("noten", {}).get(art, {}) or {}).get(nid, "")))
        r = run._r
        for k in list(r):
            r.remove(k)
        rpr = OxmlElement("w:rPr")
        rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "FootnoteReference" if art == "f" else "EndnoteReference")
        rpr.append(rs)
        r.append(rpr)
        ref = OxmlElement("w:footnoteReference" if art == "f" else "w:endnoteReference")
        ref.set(qn("w:id"), str(self._note_ids[key]))
        r.append(ref)

    def _zusatz_anwenden(self, run, cf: QTextCharFormat) -> None:
        """Kommentarbereich merken, nachverfolgte Änderung als w:ins / w:del um den Lauf legen."""
        if cf.hasProperty(P_KOMMENTAR):
            self._komm_runs.setdefault(int(cf.property(P_KOMMENTAR)), []).append(run)
        if cf.hasProperty(P_AENDERUNG):
            art = cf.stringProperty(P_AENDERUNG)
            if art in ("ins", "del"):
                self._aend_id += 1
                w = OxmlElement("w:ins" if art == "ins" else "w:del")
                w.set(qn("w:id"), str(1000 + self._aend_id))
                w.set(qn("w:author"), cf.stringProperty(P_AUTOR) or "Unbekannt")
                if cf.stringProperty(P_DATUM):
                    w.set(qn("w:date"), cf.stringProperty(P_DATUM))
                run._r.addprevious(w)
                w.append(run._r)
                if art == "del":
                    for t in run._r.findall(qn("w:t")):
                        t.tag = qn("w:delText")

    def _steuerelement(self, form: str, wert: str):
        """w:sdt für ein Formularfeld: Art (Text, Liste, Datum, Kontrollkästchen), Name als Titel, unsere Daten im Tag."""
        from lunivo_writer.editor import formular as FO
        m = FO.meta_lesen(form) or {"art": "text", "name": ""}
        W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
        sdt = OxmlElement("w:sdt")
        pr = OxmlElement("w:sdtPr")
        sdt.append(pr)
        if m.get("name"):
            al = OxmlElement("w:alias"); al.set(qn("w:val"), str(m["name"])); pr.append(al)
        tag = OxmlElement("w:tag"); tag.set(qn("w:val"), "lunivo:" + form); pr.append(tag)
        if m.get("leer"):
            pr.append(OxmlElement("w:showingPlcHdr"))
        art = m["art"]
        if art == "text":
            pr.append(OxmlElement("w:text"))
        elif art == "auswahl":
            dd = OxmlElement("w:dropDownList")
            for o in m.get("optionen") or []:
                li = OxmlElement("w:listItem"); li.set(qn("w:displayText"), str(o)); li.set(qn("w:value"), str(o)); dd.append(li)
            pr.append(dd)
        elif art == "datum":
            da = OxmlElement("w:date")
            fm = OxmlElement("w:dateFormat"); fm.set(qn("w:val"), FO.DATUMSFORMAT); da.append(fm)
            lid = OxmlElement("w:lid"); lid.set(qn("w:val"), "de-DE"); da.append(lid)
            pr.append(da)
        elif art == "kasten":
            from lxml import etree
            cb = etree.SubElement(pr, "{%s}checkbox" % W14)
            for nm, val in (("checked", "1" if wert.strip() == FO.KASTEN_VOLL else "0"), ("checkedState", "2612"), ("uncheckedState", "2610")):
                el = etree.SubElement(cb, "{%s}%s" % (W14, nm)); el.set("{%s}val" % W14, val)
        inhalt = OxmlElement("w:sdtContent")
        sdt.append(inhalt)
        return sdt, inhalt

    def _make_run(self, parent, para):
        if parent is para:
            return para.add_run()
        from docx.text.run import Run
        r = OxmlElement("w:r")
        parent.append(r)
        return Run(r, para)

    def _hyperlink(self, para, href: str):
        h = OxmlElement("w:hyperlink")
        if href.startswith("#"):
            h.set(qn("w:anchor"), href[1:])
        else:
            rid = para.part.relate_to(href, RT.HYPERLINK, is_external=True)
            h.set(qn("r:id"), rid)
        para._p.append(h)
        return h

    def _format_run(self, run, cf: QTextCharFormat, in_head: bool, is_link: bool):
        f = run.font
        if _has(cf, QTextFormat.FontFamilies) and cf.fontFamilies():
            fam = cf.fontFamilies()[0]
            f.name = fam
            rpr = run._r.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            for a in ("w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), fam)
        elif _has(cf, QTextFormat.FontFamily):
            f.name = cf.stringProperty(QTextFormat.FontFamily)
        if _has(cf, QTextFormat.FontPointSize) and cf.fontPointSize() > 0:
            f.size = Pt(cf.fontPointSize())
        if _has(cf, QTextFormat.FontWeight):
            f.bold = cf.fontWeight() >= QFont.DemiBold
        if _has(cf, QTextFormat.FontItalic):
            f.italic = cf.fontItalic()
        ustil = cf.underlineStyle()
        if ustil != QTextCharFormat.NoUnderline and not is_link:
            from docx.enum.text import WD_UNDERLINE
            f.underline = {QTextCharFormat.DashUnderline: WD_UNDERLINE.DASH, QTextCharFormat.DotLine: WD_UNDERLINE.DOTTED,
                           QTextCharFormat.DashDotLine: WD_UNDERLINE.DOT_DASH,
                           QTextCharFormat.WaveUnderline: WD_UNDERLINE.WAVY}.get(ustil, True)
        if cf.fontStrikeOut():
            f.strike = True
        va = cf.verticalAlignment()
        if va == QTextCharFormat.AlignSuperScript:
            f.superscript = True
        elif va == QTextCharFormat.AlignSubScript:
            f.subscript = True
        fg = cf.foreground()
        if fg.style() != Qt.NoBrush and not is_link:
            c = fg.color()
            if c.isValid() and c.name().lower() != "#000000":
                f.color.rgb = RGBColor(c.red(), c.green(), c.blue())
        bg = cf.background()
        if bg.style() != Qt.NoBrush:
            c = bg.color()
            if c.isValid():
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), c.name()[1:].upper())
                run._r.get_or_add_rPr().append(shd)
        if cf.textOutline().style() != Qt.NoPen:
            f.outline = True
        if cf.textOutline().style() != Qt.NoPen or (cf.hasProperty(QTextFormat.ForegroundBrush) and cf.foreground().color().alpha() < 255):
            from lunivo_writer.editor import texteffekt
            texteffekt.in_rpr(run._r.get_or_add_rPr(), cf)
        if cf.hasProperty(P_VERSTECKT):
            f.hidden = True
        ls = cf.fontLetterSpacing()
        if ls and cf.fontLetterSpacingType() == QFont.AbsoluteSpacing:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(int(round(ls * 0.75 * 20))))             # px → pt → Zwanzigstelpunkt
            run._r.get_or_add_rPr().insert_element_before(
                sp, "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
                "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
        cap = cf.fontCapitalization()
        if cap == QFont.AllUppercase:
            f.all_caps = True
        elif cap == QFont.SmallCaps:
            f.small_caps = True
        if is_link:
            run._r.get_or_add_rPr().append(_rstyle_hyperlink(self.d))

    # Kompatibilitätsoptionen (Optionen ▸ Allgemein und Speichern): als echte Einstellungen in settings.xml; Reihenfolge nach ECMA-376
    KOMPAT_VOR_FE = ("ulTrailSpace", "footnoteLayoutLikeWW8", "adjustLineHeightInTable")
    KOMPAT_NACH_FE = ("doNotUseIndentAsNumberingTabStop", "splitPgBreakAndParaMark")

    def _kompatibilitaet(self) -> None:
        gewaehlt = set(self.meta.get("kompat") or [])
        einst = self.d.settings.element
        if "bilder_komprimieren" in self.meta:
            wort = einst.find(qn("w:doNotAutoCompressPictures"))
            if self.meta["bilder_komprimieren"] and wort is not None:
                einst.remove(wort)                                    # Word darf Bilder beim Speichern selbst komprimieren
        if not gewaehlt:
            return
        compat = einst.find(qn("w:compat"))
        if compat is None:
            return
        for name in self.KOMPAT_VOR_FE[::-1]:
            if name in gewaehlt and compat.find(qn(f"w:{name}")) is None:
                compat.insert(0, OxmlElement(f"w:{name}"))
        fe = compat.find(qn("w:useFELayout"))
        stelle = list(compat).index(fe) + 1 if fe is not None else 0
        for name in self.KOMPAT_NACH_FE:
            if name in gewaehlt and compat.find(qn(f"w:{name}")) is None:
                compat.insert(stelle, OxmlElement(f"w:{name}"))
                stelle += 1

    def _image(self, imf, para, seite: str | None = None):
        name = imf.name()
        img = self.q.resource(QTextDocument.ImageResource, QUrl(name))
        if not isinstance(img, QImage) or img.isNull():
            img = QImage(name)
        if img.isNull():
            self.report.warn(f"Bild „{os.path.basename(name)}“ nicht lesbar und nicht gespeichert.")
            return
        if self.meta.get("bilder_komprimieren") and max(img.width(), img.height()) > 1600:        # „Bilder nicht komprimieren“ ist aus
            img = img.scaled(1600, 1600, Qt.KeepAspectRatio, Qt.SmoothTransformation)
        buf = QBuffer()
        buf.open(QIODevice.WriteOnly)
        img.save(buf, "PNG")
        data = bytes(buf.data())
        w = imf.width() or img.width()
        h = imf.height() or img.height()
        run = para.add_run()
        run.add_picture(io.BytesIO(data), width=Emu(int(w * EMU_PER_PX)), height=Emu(int(h * EMU_PER_PX)))
        alt = imf.property(STYLE_PROP + 1) if imf.hasProperty(STYLE_PROP + 1) else ""
        if alt:
            for dp in run._r.xpath(".//wp:docPr"):
                dp.set("descr", str(alt))
        if imf.hasProperty(P_OBJEKT):                       # Daten zum Wiederbearbeiten (Form, Diagramm, SmartArt …) im Titel des Bildes
            for dp in run._r.xpath(".//wp:docPr"):
                dp.set("title", "lunivo:" + str(imf.property(P_OBJEKT)))
        if seite:
            self._verankern(run, "right" if seite == "rechts" else "left")

    def _verankern(self, run, ausrichtung: str) -> None:
        """Macht aus dem Bild im Text (wp:inline) ein umflossenes Bild (wp:anchor, Umbruch „Quadrat“, links/rechts)."""
        for inline in run._r.xpath(".//wp:inline"):
            anker = OxmlElement("wp:anchor")
            for k, v in (("distT", "0"), ("distB", "0"), ("distL", "114300"), ("distR", "114300"), ("simplePos", "0"),
                         ("relativeHeight", "251658240"), ("behindDoc", "0"), ("locked", "0"), ("layoutInCell", "1"),
                         ("allowOverlap", "0")):
                anker.set(k, v)
            sp = OxmlElement("wp:simplePos"); sp.set("x", "0"); sp.set("y", "0")
            ph = OxmlElement("wp:positionH"); ph.set("relativeFrom", "margin")
            al = OxmlElement("wp:align"); al.text = ausrichtung; ph.append(al)
            pv = OxmlElement("wp:positionV"); pv.set("relativeFrom", "paragraph")
            po = OxmlElement("wp:posOffset"); po.text = "0"; pv.append(po)
            ee = OxmlElement("wp:effectExtent")
            for k in ("l", "t", "r", "b"):
                ee.set(k, "0")
            wr = OxmlElement("wp:wrapSquare"); wr.set("wrapText", "bothSides")
            extent = inline.find(qn("wp:extent"))
            vorn = [sp, ph, pv] + ([extent] if extent is not None else [])
            for e in vorn + [ee, wr]:
                anker.append(e)
            for k in list(inline):
                if k is not extent and k.tag != qn("wp:effectExtent"):
                    anker.append(k)
            inline.getparent().replace(inline, anker)

    # ---- Tabellen -------------------------------------------------------
    def write_table(self, qt: QTextTable, container):
        rows, cols = qt.rows(), qt.columns()
        t = container.add_table(rows, cols)
        if qt.format().border() > 0:
            t.style = self.d.styles["Table Grid"]
        done = set()
        for r in range(rows):
            for c in range(cols):
                cell = qt.cellAt(r, c)
                key = (cell.row(), cell.column())
                if key in done:
                    continue
                done.add(key)
                rs, cs = cell.rowSpan(), cell.columnSpan()
                dcell = t.cell(r, c)
                if rs > 1 or cs > 1:
                    dcell = dcell.merge(t.cell(r + rs - 1, c + cs - 1))
                self._zelle_format(dcell, cell.format().toTableCellFormat())
                box = _CellBox(dcell)
                # Zelleninhalt: Rahmen der Zelle durchlaufen
                cit = cell.begin()
                first = True
                while not cit.atEnd():
                    blk = cit.currentBlock()
                    if blk.isValid():
                        para = box.first_paragraph() if first else box.add_paragraph()
                        first = False
                        self.write_block(blk, para)
                    cit += 1
        if isinstance(container, _Body):
            pass

    @staticmethod
    def _zelle_format(dcell, cf) -> None:
        """Zellenschattierung (w:shd) und senkrechte Ausrichtung (w:vAlign) der Qt-Zelle übernehmen."""
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT as V
        br = cf.background()
        if br.style() != Qt.NoBrush and br.color().alpha() > 0:
            tcpr = dcell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), br.color().name()[1:].upper())
            for alt in tcpr.findall(qn("w:shd")):
                tcpr.remove(alt)
            tcpr.append(shd)
        va = {QTextCharFormat.AlignMiddle: V.CENTER, QTextCharFormat.AlignBottom: V.BOTTOM}.get(cf.verticalAlignment())
        if va is not None:
            dcell.vertical_alignment = va

    # ---- Hauptlauf --------------------------------------------------------
    def run(self) -> bytes:
        self.write_defaults()
        self._sprache()
        self.write_page()
        self.write_frame(self.q.rootFrame(), _Body(self.d))
        self._trim_trailing_empty()
        self._noten_schreiben()
        self._kommentare_schreiben()
        self._quellen_schreiben()
        self._kompatibilitaet()
        cp = self.d.core_properties
        cp.title = self.meta.get("title", "") or ""
        cp.author = self.meta.get("author", "") or "Lunivo Writer"
        out = io.BytesIO()
        self.d.save(out)
        return out.getvalue()

    def _kommentare_schreiben(self) -> None:
        daten = self.meta.get("kommentare", {}) or {}
        for kid in sorted(self._komm_runs):
            runs = self._komm_runs[kid]
            k = daten.get(kid) or daten.get(str(kid)) or {}
            try:
                c = self.d.add_comment(runs=[runs[0], runs[-1]] if len(runs) > 1 else runs[0], text=k.get("text", ""),
                                       author=k.get("autor", "") or "Unbekannt", initials=k.get("initialen", "") or "")
                if k.get("datum"):
                    c._comment_elm.set(qn("w:date"), k["datum"])
            except Exception as ex:
                self.report.warn(f"Ein Kommentar konnte nicht gespeichert werden: {ex}")

    def _quellen_schreiben(self) -> None:
        quellen = self.meta.get("quellen") or []
        if not quellen:
            return
        import uuid
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        from lxml import etree
        B = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"
        w = etree.Element(f"{{{B}}}Sources", nsmap={"b": B, None: B},
                          SelectedStyle="\\APA.XSL", StyleName=self.meta.get("zitierweise", "APA"), Version="6")
        for q in quellen:
            e = etree.SubElement(w, f"{{{B}}}Source")
            def sub(n, v, parent=e):
                x = etree.SubElement(parent, f"{{{B}}}{n}")
                x.text = v or ""
                return x
            sub("Tag", q.get("tag", ""))
            sub("SourceType", {"Buch": "Book", "Aufsatz": "JournalArticle", "Webseite": "InternetSite", "Bericht": "Report"}.get(q.get("typ"), "Misc"))
            a1 = etree.SubElement(etree.SubElement(e, f"{{{B}}}Author"), f"{{{B}}}Author")
            nl = etree.SubElement(a1, f"{{{B}}}NameList")
            for person in [x.strip() for x in (q.get("autor") or "").split(";") if x.strip()]:
                pe = etree.SubElement(nl, f"{{{B}}}Person")
                last, _, first = person.partition(",")
                sub("Last", last.strip(), pe)
                if first.strip():
                    sub("First", first.strip(), pe)
            sub("Title", q.get("titel", "")); sub("Year", q.get("jahr", "")); sub("City", q.get("ort", ""))
            sub("Publisher", q.get("verlag", "")); sub("URL", q.get("url", ""))
        blob = etree.tostring(w, xml_declaration=True, encoding="UTF-8", standalone=True)
        pkg = self.d.part.package
        vorhanden = {str(x.partname) for x in pkg.iter_parts()}
        n = 1
        while f"/customXml/item{n}.xml" in vorhanden or f"/customXml/itemProps{n}.xml" in vorhanden:
            n += 1
        item = Part(PackURI(f"/customXml/item{n}.xml"), "application/xml", blob, pkg)
        guid = "{" + str(uuid.uuid4()).upper() + "}"
        props = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 f'<ds:datastoreItem ds:itemID="{guid}" xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml">'
                 f'<ds:schemaRefs><ds:schemaRef ds:uri="{B}"/></ds:schemaRefs></ds:datastoreItem>').encode("utf-8")
        pp = Part(PackURI(f"/customXml/itemProps{n}.xml"),
                  "application/vnd.openxmlformats-officedocument.customXmlProperties+xml", props, pkg)
        item.relate_to(pp, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps")
        self.d.part.relate_to(item, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml")

    def _noten_schreiben(self) -> None:
        from docx.enum.style import WD_STYLE_TYPE
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        from lxml import etree
        W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        for art, wurzel, tag, ref, rel, ct, stil_t, stil_r in (
                ("f", "footnotes", "footnote", "footnoteRef", "footnotes",
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", "Footnote Text", "Footnote Reference"),
                ("e", "endnotes", "endnote", "endnoteRef", "endnotes",
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml", "Endnote Text", "Endnote Reference")):
            if not self._noten[art]:
                continue
            names = [s.name for s in self.d.styles]
            if stil_t not in names:
                st = self.d.styles.add_style(stil_t, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = self.d.styles["Normal"]
                st.font.size = Pt(10)
            if stil_r not in names:
                st = self.d.styles.add_style(stil_r, WD_STYLE_TYPE.CHARACTER)
                st.font.superscript = True
            root = etree.Element(f"{{{W}}}{wurzel}", nsmap={"w": W})

            def trenner(typ, nid, kind):
                n = etree.SubElement(root, f"{{{W}}}{tag}", {f"{{{W}}}type": typ, f"{{{W}}}id": str(nid)})
                p = etree.SubElement(n, f"{{{W}}}p")
                r = etree.SubElement(p, f"{{{W}}}r")
                etree.SubElement(r, f"{{{W}}}{kind}")
            trenner("separator", -1, "separator")
            trenner("continuationSeparator", 0, "continuationSeparator")
            for nid, text in self._noten[art]:
                n = etree.SubElement(root, f"{{{W}}}{tag}", {f"{{{W}}}id": str(nid)})
                for i, zeile in enumerate((text or "").split("\n")):
                    p = etree.SubElement(n, f"{{{W}}}p")
                    ppr = etree.SubElement(p, f"{{{W}}}pPr")
                    etree.SubElement(ppr, f"{{{W}}}pStyle", {f"{{{W}}}val": stil_t.replace(" ", "")})
                    if i == 0:
                        r = etree.SubElement(p, f"{{{W}}}r")
                        rpr = etree.SubElement(r, f"{{{W}}}rPr")
                        etree.SubElement(rpr, f"{{{W}}}rStyle", {f"{{{W}}}val": stil_r.replace(" ", "")})
                        etree.SubElement(r, f"{{{W}}}{ref}")
                    r2 = etree.SubElement(p, f"{{{W}}}r")
                    t = etree.SubElement(r2, f"{{{W}}}t")
                    t.text = (" " if i == 0 else "") + zeile
                    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            part = Part(PackURI(f"/word/{wurzel}.xml"), ct, blob, self.d.part.package)
            self.d.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/" + rel)

    def _trim_trailing_empty(self):
        body = self.d.element.body
        kids = [k for k in body if k.tag != qn("w:sectPr")]
        # Word verlangt nach einer Tabelle einen Absatz; leere Endabsätze sonst entfernen
        while len(kids) > 1 and kids[-1].tag == qn("w:p") and not "".join(kids[-1].itertext()).strip() \
                and kids[-2].tag == qn("w:p") and not kids[-1].xpath(".//w:drawing"):
            body.remove(kids[-1])
            kids.pop()




def _split(text: str):
    buf = ""
    for ch in text:
        if ch in ("\t", " ", "￼", "\n"):
            if buf:
                yield buf, "text"
                buf = ""
            yield "", {"\t": "tab", " ": "br", "\n": "br", "￼": "img"}[ch]
        else:
            buf += ch
    if buf:
        yield buf, "text"


def _rstyle_hyperlink(d):
    el = OxmlElement("w:rStyle")
    el.set(qn("w:val"), "Hyperlink")
    try:
        d.styles["Hyperlink"]
    except KeyError:
        st = d.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        st.font.underline = True
        st.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    return el


class _Body:
    def __init__(self, d):
        self.d = d

    def add_paragraph(self):
        return self.d.add_paragraph()

    def add_table(self, r, c):
        return self.d.add_table(rows=r, cols=c)


class _CellBox:
    def __init__(self, cell):
        self.cell = cell

    def first_paragraph(self):
        return self.cell.paragraphs[0]

    def add_paragraph(self):
        return self.cell.add_paragraph()

    def add_table(self, r, c):
        return self.cell.add_table(r, c)


def write_docx(path: str, qdoc: QTextDocument, page: PageSettings, meta: Optional[dict] = None) -> IOReport:
    """Schreibt atomar (Temp-Datei, dann umbenennen). Wirft DocumentIOError."""
    w = _Writer(qdoc, page, meta or {})
    try:
        data = w.run()
    except Exception as e:
        raise DocumentIOError(f"Das Dokument konnte nicht als DOCX aufgebaut werden: {e}") from e
    folder = os.path.dirname(os.path.abspath(path))
    try:
        fd, tmp = tempfile.mkstemp(dir=folder, suffix=".tmp")
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        os.replace(tmp, path)
    except OSError as e:
        raise DocumentIOError(f"Speichern nicht möglich: {e.strerror or e}") from e
    return w.report
