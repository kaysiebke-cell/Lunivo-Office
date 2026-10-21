"""DOCX → QTextDocument (python-docx).

Gelesen werden: Absätze mit Formatvorlage, Überschriften, Zeichenformat
(Schrift, Größe, Fett, Kursiv, Unterstrichen, Durchgestrichen, Hoch/Tief,
Farbe, Hervorhebung), Ausrichtung, Einzüge, Abstände, Zeilenabstand,
Aufzählungen und Nummerierungen, Seitenumbrüche, Tabellen (mit verbundenen
Zellen), Bilder im Text, Hyperlinks, Seitengröße, Ränder, Kopf- und Fußzeile.

Alles, was nicht gelesen wird, landet im IOReport — nichts verschwindet still.
"""
from __future__ import annotations

import io
import zipfile
from typing import Optional

import docx
from docx.document import Document as _Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.run import Run
from docx.text.paragraph import Paragraph
from PySide6.QtCore import QUrl
from PySide6.QtGui import (
    QBrush, QColor, QFont, QImage, QTextBlockFormat, QTextCharFormat,
    QTextCursor, QTextDocument, QTextFormat, QTextImageFormat,
    QTextLength, QTextListFormat, QTextTableFormat,
)

from lunivo_writer.models.page import HeaderFooter, PageSettings, px_to_mm
from lunivo_writer.models.report import DocumentIOError, IOReport

LH_PROPORTIONAL, LH_FIXED = 1, 2                # QTextBlockFormat.LineHeightTypes
STYLE_PROP = QTextFormat.UserProperty + 1       # Name der Formatvorlage am Block
from lunivo_writer.editor.zusatz import B_SPALTE, P_OBJEKT, P_VERSTECKT  # noqa: E402
from lunivo_writer.editor.formular import P_FORM  # noqa: E402
from lunivo_writer.editor.linie import Linie, art_aus_docx, docx_wert
from lunivo_writer.editor import muster
from lunivo_writer.editor.zusatz import B_ART, B_CAP, B_MUSTER, B_RAHMEN, B_TOC, BESCHRIFTUNGEN, P_AENDERUNG, P_AUTOR, P_DATUM, P_FELD, P_INDEX, P_KOMMENTAR, P_NOTE, P_ZITAT, leere_meta
EMU_PER_PX = 9525
PT_PER_PX = 0.75

HEADING_DEFAULT_PT = {1: 20, 2: 16, 3: 14, 4: 12, 5: 11, 6: 11, 7: 11, 8: 11, 9: 11}
NAMED_STYLE_DEFAULT = {  # Name → (Punktgröße, fett, kursiv)
    "Title": (26, False, False),
    "Subtitle": (15, False, True),
    "Quote": (None, False, True),
}

from PySide6.QtCore import Qt  # noqa: E402
_QT_ALIGN = {
    WD_ALIGN_PARAGRAPH.LEFT: Qt.AlignLeft,
    WD_ALIGN_PARAGRAPH.CENTER: Qt.AlignHCenter,
    WD_ALIGN_PARAGRAPH.RIGHT: Qt.AlignRight,
    WD_ALIGN_PARAGRAPH.JUSTIFY: Qt.AlignJustify,
}

# Was beim Lesen nicht übernommen wird → Hinweis
_MC_ALT = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
_MC_CHOICE = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice"

_UNSUPPORTED = [
    ("//w:txbxContent", "Textfelder werden nicht dargestellt und gehen beim Speichern verloren."),
    ("//m:oMath", "Formeln werden nicht dargestellt und gehen beim Speichern verloren."),
    ("//w:pict | //wps:wsp[not(.//a:blip)] | //wpg:wgp", "Formen und gezeichnete Objekte werden nicht dargestellt."),
    ("//w:drawing[.//wp:anchor[.//wp:wrapTopAndBottom]]",
     "Bilder mit „Oben und unten“-Umbruch werden als Bild im Text eingefügt."),
]


from docx.oxml.ns import nsmap as _NS  # noqa: E402


def _xp(el, ausdruck: str):
    """XPath mit den Namensräumen von python-docx — auch für Elemente, deren Klasse kein eigenes xpath mit Namensräumen hat."""
    from lxml import etree
    return etree._Element.xpath(el, ausdruck, namespaces=_NS)


def _emu_to_px(v) -> float:
    return float(v) / EMU_PER_PX if v is not None else 0.0


def _pt_to_px(v) -> float:
    return float(v) / PT_PER_PX


def _style_chain(style):
    seen = set()
    while style is not None and id(style) not in seen:
        seen.add(id(style))
        yield style
        style = style.base_style


def _doc_defaults(document: _Document) -> tuple[Optional[str], Optional[float]]:
    name, size = None, None
    styles = document.styles.element
    rpr = styles.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if rpr is not None:
        f = rpr.find(qn("w:rFonts"))
        if f is not None:
            name = f.get(qn("w:ascii"))
        s = rpr.find(qn("w:sz"))
        if s is not None:
            size = int(s.get(qn("w:val"))) / 2.0
    return name, size


class _Reader:
    def __init__(self, path: str):
        self.path = path
        self.report = IOReport()
        try:
            self.d = docx.Document(path)
        except (zipfile.BadZipFile, KeyError, ValueError) as e:
            raise DocumentIOError(
                "Die Datei ist kein gültiges DOCX-Dokument oder beschädigt.") from e
        except Exception as e:  # python-docx wirft viele Typen bei kaputten Paketen
            raise DocumentIOError(f"Die Datei konnte nicht gelesen werden: {e}") from e
        self.doc = QTextDocument()
        self.doc.setDocumentMargin(0)
        self.cur = QTextCursor(self.doc)
        self.fresh = True                 # aktueller Block ist leer und noch frei
        self.last_list = None             # (key, QTextList)
        self.pending_break = False
        self._ab_zaehler = 0
        self.neue_vorlagen: list[str] = []
        self._erste_seite = None
        self.deferred_break = False       # Seitenumbruch im Absatz, danach folgt noch nichts
        self.img_n = 0
        self.zus = leere_meta()
        self._komm_offen: list[int] = []
        self._lz_offen: dict[str, str] = {}
        self._lz_pos: dict[str, int] = {}
        self._lz_wartend: list[str] = []
        self._aend = None                 # (art, autor, datum) innerhalb w:ins / w:del
        self._form = None                 # JSON-Text des Formularfelds innerhalb w:sdt
        self._note_n = {"f": 0, "e": 0}
        self._feld = None                 # {"instr", "phase": "i"|"r"} innerhalb eines Feldes
        self._letzter = None              # (von, bis) des zuletzt eingefügten Textstücks (für XE-Felder)
        self._laden_zusatz()
        self._stile_id = {st.style_id: st for st in self.d.styles}          # Zwischenspeicher: python-docx sucht sonst bei jedem Zugriff per XPath
        self._standard_absatz = self._standard_stil(WD_STYLE_TYPE.PARAGRAPH)
        self._standard_zeichen = self._standard_stil(WD_STYLE_TYPE.CHARACTER)
        self._basis: dict[int, object] = {}
        self._ketten: dict[int, list] = {}
        self._attr_cache: dict[tuple, object] = {}
        self.def_font, self.def_size = _doc_defaults(self.d)
        normal = self.d.styles["Normal"] if "Normal" in [s.name for s in self.d.styles] else None
        self.normal = normal

    def _laden_zusatz(self) -> None:
        """Kommentare (python-docx) und Fuß-/Endnoten (Teile footnotes.xml / endnotes.xml) einlesen."""
        try:
            for c in self.d.comments:
                self.zus["kommentare"][int(c.comment_id)] = {
                    "autor": c.author or "", "initialen": c.initials or "",
                    "datum": c.timestamp.isoformat() if c.timestamp else "",
                    "text": "\n".join(p.text for p in c.paragraphs)}
        except Exception:
            self.report.warn("Kommentare konnten nicht gelesen werden.")
        self._quellen_lesen()
        for art, rel in (("f", "footnotes"), ("e", "endnotes")):
            for r in self.d.part.rels.values():
                if r.reltype.endswith("/" + rel) and not r.is_external:
                    from lxml import etree
                    wurzel = etree.fromstring(r.target_part.blob)
                    tag = qn("w:footnote") if art == "f" else qn("w:endnote")
                    for n in wurzel.iterchildren(tag):
                        if n.get(qn("w:type")) in ("separator", "continuationSeparator", "continuationNotice"):
                            continue
                        zeilen = []
                        for p in n.iterchildren(qn("w:p")):
                            zeilen.append("".join(t.text or "" for t in p.iter(qn("w:t"))).strip())
                        self.zus["noten"][art][int(n.get(qn("w:id")))] = "\n".join(zeilen)

    def _quellen_lesen(self) -> None:
        """Literaturquellen aus dem Word-Bibliografieteil (customXml, Namensraum …/bibliography)."""
        from lxml import etree
        B = "{http://schemas.openxmlformats.org/officeDocument/2006/bibliography}"
        for r in self.d.part.rels.values():
            if r.is_external or not r.reltype.endswith("/customXml"):
                continue
            try:
                w = etree.fromstring(r.target_part.blob)
            except Exception:
                continue
            if w.tag != B + "Sources":
                continue
            self.zus["zitierweise"] = (w.get("StyleName") or "APA")
            for q in w.iterchildren(B + "Source"):
                def t(n):
                    e = q.find(B + n)
                    return (e.text or "") if e is not None else ""
                personen = []
                for pe in q.iter(B + "Person"):
                    last = (pe.findtext(B + "Last") or "").strip(); first = (pe.findtext(B + "First") or "").strip()
                    personen.append(f"{last}, {first}" if first else last)
                corp = q.find(f"{B}Author/{B}Author/{B}Corporate")
                if corp is not None and corp.text:
                    personen.append(corp.text)
                typ = {"Book": "Buch", "JournalArticle": "Aufsatz", "InternetSite": "Webseite", "Report": "Bericht"}.get(t("SourceType"), "Sonstiges")
                self.zus["quellen"].append({"tag": t("Tag"), "typ": typ, "autor": "; ".join(personen), "titel": t("Title"),
                                            "jahr": t("Year"), "verlag": t("Publisher"), "ort": t("City"), "url": t("URL")})

    # ---- Auflösung der Vererbung -------------------------------------
    def _standard_stil(self, art):
        try:
            return self.d.styles.default(art)
        except Exception:
            return None

    def _ps(self, para):
        """Absatzvorlage (mit Zwischenspeicher)."""
        ppr = para._p.pPr
        sid = ppr.pStyle.val if ppr is not None and ppr.pStyle is not None else None
        return self._stile_id.get(sid) or self._standard_absatz

    def _rs(self, run):
        rpr = run._r.rPr
        sid = rpr.rStyle.val if rpr is not None and rpr.rStyle is not None else None
        return self._stile_id.get(sid) if sid else None

    def _kette(self, style):
        k = self._ketten.get(id(style.element))
        if k is None:
            k, gesehen, s = [], set(), style
            while s is not None and id(s.element) not in gesehen:
                gesehen.add(id(s.element)); k.append(s)
                b = self._basis.get(id(s.element), False)
                if b is False:
                    b = s.base_style
                    self._basis[id(s.element)] = b
                s = b
            self._ketten[id(style.element)] = k
        return k

    def _style_attr(self, style, getter):
        for s in self._kette(style):
            v = getter(s)
            if v is not None:
                return v
        return None

    def _stil_font(self, style, attr):
        key = (id(style.element), attr)
        if key not in self._attr_cache:
            self._attr_cache[key] = self._style_attr(style, lambda s: getattr(s.font, attr) if attr != "name" else s.font.name)
        return self._attr_cache[key]

    def _font_name(self, run, para) -> str:
        v = run.font.name
        rs = self._rs(run) if hasattr(run, "_r") else None
        if v is None and rs is not None:
            v = self._stil_font(rs, "name")
        if v is None and para is not None:
            v = self._stil_font(self._ps(para), "name")
        if v is None and self.normal is not None:
            v = self.normal.font.name
        return v or self.def_font or "Calibri"

    def _font_size_pt(self, run, para) -> float:
        v = run.font.size
        rs = self._rs(run) if hasattr(run, "_r") else None
        if v is None and rs is not None:
            v = self._stil_font(rs, "size")
        if v is None and para is not None:
            v = self._stil_font(self._ps(para), "size")
        if v is not None:
            return v.pt
        if para is not None:
            lvl = self._heading_level(para)
            if lvl:
                return HEADING_DEFAULT_PT[lvl]
            nm = NAMED_STYLE_DEFAULT.get(self._ps(para).name)
            if nm and nm[0]:
                return nm[0]
        if self.normal is not None and self.normal.font.size is not None:
            return self.normal.font.size.pt
        return self.def_size or 11.0

    def _flag(self, run, para, attr):
        v = getattr(run.font, attr)
        rs = self._rs(run) if hasattr(run, "_r") else None
        if v is None and rs is not None:
            v = self._stil_font(rs, attr)
        if v is None and para is not None:
            v = self._stil_font(self._ps(para), attr)
        return v

    def _heading_level(self, para) -> int:
        st = self._ps(para)
        n = st.name if st is not None else ""
        if n.startswith("Heading ") and n[8:].isdigit():
            return max(1, min(9, int(n[8:])))
        return 0

    # ---- Zeichenformat -----------------------------------------------
    def _char_format(self, run, para) -> QTextCharFormat:
        f = QTextCharFormat()
        f.setFontFamilies([self._font_name(run, para)])
        f.setFontPointSize(self._font_size_pt(run, para))
        lvl = self._heading_level(para) if para is not None else 0
        named = NAMED_STYLE_DEFAULT.get(self._ps(para).name) if para is not None else None
        bold = self._flag(run, para, "bold")
        if bold is None and lvl:
            bold = True
        f.setFontWeight(QFont.Bold if bold else QFont.Normal)
        ital = self._flag(run, para, "italic")
        if ital is None and named:
            ital = named[2]
        f.setFontItalic(bool(ital))
        ul = self._flag(run, para, "underline")
        f.setFontUnderline(bool(ul))
        raw = run.font.underline
        if raw not in (None, True, False):
            from docx.enum.text import WD_UNDERLINE as U
            from lunivo_writer.ui.baukasten.text_format import unter_format
            art = {U.WORDS: "woerter", U.DOUBLE: "doppelt", U.THICK: "dick", U.DOTTED: "gepunktet", U.DOTTED_HEAVY: "dick_gepunktet", U.DASH: "gestrichelt",
                   U.DASH_HEAVY: "dick_gestrichelt", U.DASH_LONG: "lang", U.DASH_LONG_HEAVY: "dick_lang", U.DOT_DASH: "strichpunkt",
                   U.DOT_DASH_HEAVY: "dick_strichpunkt", U.DOT_DOT_DASH: "strichpunktpunkt", U.DOT_DOT_DASH_HEAVY: "dick_strichpunktpunkt",
                   U.WAVY: "wellig", U.WAVY_HEAVY: "dick_wellig", U.WAVY_DOUBLE: "doppelwellig"}.get(raw)
            if art:
                f.merge(unter_format(art))
        if run.font.outline:
            from PySide6.QtGui import QPen
            f.setTextOutline(QPen(QColor("#000000")))
        if run.font.hidden:
            f.setProperty(P_VERSTECKT, 1)
        rpr = run._r.rPr
        sp = rpr.find(qn("w:spacing")) if rpr is not None else None
        if sp is not None and sp.get(qn("w:val")):
            f.setFontLetterSpacingType(QFont.AbsoluteSpacing)
            f.setFontLetterSpacing(int(sp.get(qn("w:val"))) / 20 / 0.75)
        f.setFontStrikeOut(bool(self._flag(run, para, "strike")))
        if self._flag(run, para, "superscript"):
            f.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
        elif self._flag(run, para, "subscript"):
            f.setVerticalAlignment(QTextCharFormat.AlignSubScript)
        col = run.font.color
        if col is not None and col.rgb is not None:
            f.setForeground(QBrush(QColor("#" + str(col.rgb))))
        hl = self._highlight(run)
        if hl:
            f.setBackground(QBrush(QColor(hl)))
        if self._flag(run, para, "all_caps"):
            f.setFontCapitalization(QFont.AllUppercase)
        elif self._flag(run, para, "small_caps"):
            f.setFontCapitalization(QFont.SmallCaps)
        self._schrift_zusatz(run._r.rPr, f)
        from lunivo_writer.editor import texteffekt
        texteffekt.aus_rpr(run._r.rPr, f)
        return f

    def _schrift_zusatz(self, rpr, f: QTextCharFormat) -> None:
        """Schrift für asiatische und komplexe Zeichen, Betonungszeichen und doppeltes Durchstreichen (Fenster „Schriftart“)."""
        from lunivo_writer.editor.zusatz import P_BETONUNG, P_DOPPELT, P_FONT_ASIEN, P_FONT_KOMPLEX
        if rpr is None:
            return
        rf = rpr.find(qn("w:rFonts"))
        haupt = f.fontFamilies()[0] if f.fontFamilies() else ""
        if rf is not None:
            fam = [haupt]
            for attr, prop in (("w:eastAsia", P_FONT_ASIEN), ("w:cs", P_FONT_KOMPLEX)):
                n = rf.get(qn(attr))
                if n and n != haupt:
                    schnitt = ""; gr = ""
                    if attr == "w:cs":
                        bcs, ics = rpr.find(qn("w:bCs")), rpr.find(qn("w:iCs"))
                        schnitt = ("Fett " if bcs is not None and bcs.get(qn("w:val")) not in ("0", "false") else "") + \
                                  ("Kursiv" if ics is not None and ics.get(qn("w:val")) not in ("0", "false") else "")
                        sz = rpr.find(qn("w:szCs"))
                        gr = f"{int(sz.get(qn('w:val'))) / 2:g}" if sz is not None and sz.get(qn("w:val")) else ""
                    f.setProperty(prop, f"{n}|{schnitt.strip()}|{gr}")
                    fam.append(n)
            if len(fam) > 1:
                f.setFontFamilies(list(dict.fromkeys(fam)))
        w_ = rpr.find(qn("w:w"))
        if w_ is not None and w_.get(qn("w:val")) and w_.get(qn("w:val")).isdigit():
            f.setFontStretch(int(w_.get(qn("w:val"))))                                    # Skalieren
        pos = rpr.find(qn("w:position"))
        if pos is not None and pos.get(qn("w:val")) not in (None, "0"):
            try:                                                                          # Position in Halbpunkten → Prozent der Schriftgröße
                gr = f.fontPointSize() or 11.0
                f.setBaselineOffset(int(pos.get(qn("w:val"))) / 2 / gr * 100)
            except ValueError:
                pass
        sg = rpr.find(qn("w:snapToGrid"))
        if sg is not None and sg.get(qn("w:val")) in ("0", "false", "off"):
            from lunivo_writer.editor.zusatz import P_RASTER
            f.setProperty(P_RASTER, 0)
        if rpr.find(qn("w:dstrike")) is not None and rpr.find(qn("w:dstrike")).get(qn("w:val")) not in ("0", "false"):
            f.setProperty(P_DOPPELT, 1); f.setFontStrikeOut(True)
        em = rpr.find(qn("w:em"))
        if em is not None and em.get(qn("w:val")) not in (None, "none"):
            f.setProperty(P_BETONUNG, {"dot": "punkt", "circle": "kreis", "comma": "komma", "underDot": "unten"}.get(em.get(qn("w:val")), "punkt"))

    _HL = {"yellow": "#FFFF00", "green": "#00FF00", "cyan": "#00FFFF", "magenta": "#FF00FF",
           "blue": "#0000FF", "red": "#FF0000", "darkBlue": "#000080", "darkCyan": "#008080",
           "darkGreen": "#008000", "darkMagenta": "#800080", "darkRed": "#800000",
           "darkYellow": "#808000", "darkGray": "#808080", "lightGray": "#C0C0C0",
           "black": "#000000"}

    def _highlight(self, run) -> Optional[str]:
        rpr = run._r.rPr
        if rpr is None:
            return None
        h = rpr.find(qn("w:highlight"))
        if h is not None:
            return self._HL.get(h.get(qn("w:val")))
        shd = rpr.find(qn("w:shd"))
        if shd is not None:
            fill = shd.get(qn("w:fill"))
            if fill and fill.lower() not in ("auto", "ffffff"):
                return "#" + fill
        return None

    # ---- Absatzformat ------------------------------------------------
    _JC = {"left": Qt.AlignLeft, "start": Qt.AlignLeft, "right": Qt.AlignRight, "end": Qt.AlignRight, "center": Qt.AlignHCenter,
           "both": Qt.AlignJustify, "distribute": Qt.AlignJustify, "justify": Qt.AlignJustify}

    def _ausrichtung(self, para: Paragraph, st):
        """w:jc roh lesen — python-docx kennt „start“/„end“ nicht und bricht sonst ab."""
        def roh(ppr):
            if ppr is None:
                return None
            jc = ppr.find(qn("w:jc"))
            return jc.get(qn("w:val")) if jc is not None else None
        v = roh(para._p.pPr)
        s = st
        while v is None and s is not None:
            v = roh(s.element.pPr)
            s = s.base_style
        return self._JC.get(v, Qt.AlignLeft)

    def _block_format(self, para: Paragraph) -> QTextBlockFormat:
        pf = para.paragraph_format
        st = self._ps(para)
        b = QTextBlockFormat()

        def sicher(fn, *a):
            try:
                return fn(*a)
            except (ValueError, KeyError, AttributeError):        # ungewöhnliche Werte (z. B. jc="start" von LibreOffice)
                return None

        def eff(getter):
            v = sicher(getter, pf)
            if v is None and st is not None:
                v = self._style_attr(st, lambda s: sicher(getter, s.paragraph_format))
            return v

        ppr = para._p.pPr
        if ppr is not None and ppr.find(qn("w:bidi")) is not None:
            b.setLayoutDirection(Qt.RightToLeft)
        b.setAlignment(self._ausrichtung(para, st))
        left = eff(lambda p: p.left_indent)
        right = eff(lambda p: p.right_indent)
        first = eff(lambda p: p.first_line_indent)
        if left is not None:
            lpx = _emu_to_px(left)
            info = self._list_info(para)
            if info is not None and lpx > 0:                          # Qt rückt Listenabsätze selbst um 40 px je Ebene ein
                lpx = max(0.0, lpx - 40.0 * (info[1] + 1))
            b.setLeftMargin(lpx)
        if right is not None:
            b.setRightMargin(_emu_to_px(right))
        if first is not None:
            hanging_liste = first < 0 and self._list_info(para) is not None
            # Hängender Einzug einer Liste: Das Zeichen steht links vom Text, alle Zeilen beginnen bei „links“ — Qt setzt das Zeichen selbst
            # vor den Text; ein negativer Erstzeileneinzug würde die erste Zeile gegenüber den folgenden verschieben.
            b.setTextIndent(0 if hanging_liste else _emu_to_px(first))
        sb = eff(lambda p: p.space_before)
        sa = eff(lambda p: p.space_after)
        b.setTopMargin(_pt_to_px(sb.pt) if sb is not None else 0)
        b.setBottomMargin(_pt_to_px(sa.pt) if sa is not None else 0)
        sp = ppr.find(qn("w:spacing")) if ppr is not None else None
        leer = not "".join(para._p.xpath(".//w:t/text()")).strip() if hasattr(para._p, "xpath") else False
        if sp is not None and not leer:                               # „Automatisch“ (aus Webseiten): 14 pt davor/danach; leere Absätze behalten nur ihre Höhe
            wahr = ("1", "true", "on")
            if (sp.get(qn("w:beforeAutospacing")) or "").lower() in wahr:
                b.setTopMargin(_pt_to_px(14.0))
            if (sp.get(qn("w:afterAutospacing")) or "").lower() in wahr:
                b.setBottomMargin(_pt_to_px(14.0))
        ls = eff(lambda p: p.line_spacing)
        if ls is not None:
            if isinstance(ls, float) or (isinstance(ls, int) and ls < 100):
                b.setLineHeight(float(ls) * 100, LH_PROPORTIONAL)
            else:  # Length in EMU → feste Höhe
                pt = ls.pt
                b.setLineHeight(pt / PT_PER_PX, LH_FIXED)
        if eff(lambda p: p.page_break_before) or self.pending_break:
            b.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
            if self.pending_break and getattr(self, "pending_col", False):
                b.setProperty(B_SPALTE, 1)
            self.pending_break = False
            self.pending_col = False
        shd = para._p.xpath("w:pPr/w:shd")
        art = muster.schluessel_von_ooxml(shd[0].get(qn("w:val"))) if shd else None
        if shd and art and art != "loeschen":                    # Muster über Füllfarbe
            farbe = shd[0].get(qn("w:color"))
            fill = shd[0].get(qn("w:fill"))
            mf = "#" + farbe if farbe not in (None, "auto") else "#000000"
            ff = "#" + fill if fill not in (None, "auto") else None
            b.setBackground(muster.pinsel(art, mf, ff))
            b.setProperty(B_MUSTER, muster.als_text(art, mf, ff))
        elif shd and shd[0].get(qn("w:fill")) not in (None, "auto"):
            b.setBackground(QBrush(QColor("#" + shd[0].get(qn("w:fill")))))
        from lunivo_writer.editor import absatzopt
        opt = absatzopt.aus_ppr(para._p.pPr)
        if opt:
            absatzopt.setzen(b, {**absatzopt.VORGABE, **opt})
        stops = list(pf.tab_stops)
        if stops:
            from docx.enum.text import WD_TAB_ALIGNMENT
            from PySide6.QtGui import QTextOption
            art = {WD_TAB_ALIGNMENT.RIGHT: QTextOption.RightTab, WD_TAB_ALIGNMENT.CENTER: QTextOption.CenterTab,
                   WD_TAB_ALIGNMENT.DECIMAL: QTextOption.DelimiterTab}
            b.setTabPositions([QTextOption.Tab(_emu_to_px(t.position), art.get(t.alignment, QTextOption.LeftTab))
                               for t in stops])
            from docx.enum.text import WD_TAB_LEADER
            from lunivo_writer.editor import tabfuell
            lk = {WD_TAB_LEADER.DOTS: "punkt", WD_TAB_LEADER.DASHES: "strich", WD_TAB_LEADER.LINES: "linie", WD_TAB_LEADER.MIDDLE_DOT: "mitte"}
            tabfuell.setzen(b, {round(_emu_to_px(t.position), 2): lk[t.leader] for t in stops if t.leader in lk})
        lvl = self._heading_level(para)
        if lvl:
            b.setHeadingLevel(lvl)
        if st is not None and st.name in ("Title", "Subtitle", "Quote", "No Spacing"):
            b.setProperty(STYLE_PROP, st.name)
        elif st is not None and not lvl and not st.builtin and st.name != "Normal":
            self._eigene_vorlage(st, b)
        if para._p.xpath("w:pPr/w:sectPr"):                       # letzter Absatz eines Abschnitts: der folgende Abschnitt wird hier beschrieben
            from lunivo_writer.editor import abschnitte as AB
            k, secs = self._ab_zaehler, self.d.sections
            self._ab_zaehler += 1
            if k + 1 < len(secs):
                nxt = secs[k + 1]
                b.setProperty(AB.B_ABSCHNITT, AB.eintrag(self._abschnitt_beginn(nxt), self._section_page(nxt, basis=self._erste_seite)))
        return b

    # ---- Listen ------------------------------------------------------
    def _list_info(self, para: Paragraph):
        """→ (numId, ilvl, 'bullet'|'decimal'|…) oder None."""
        ppr = para._p.pPr
        numpr = ppr.numPr if ppr is not None else None
        num_id = ilvl = None
        if numpr is not None:
            num_id = numpr.numId.val if numpr.numId is not None else None
            ilvl = numpr.ilvl.val if numpr.ilvl is not None else 0
        if num_id is None and self._ps(para) is not None:
            for s in self._kette(self._ps(para)):
                sp = s.element.pPr
                if sp is not None and sp.numPr is not None and sp.numPr.numId is not None:
                    num_id = sp.numPr.numId.val
                    ilvl = sp.numPr.ilvl.val if sp.numPr.ilvl is not None else 0
                    break
        if not num_id:
            return None
        fmt = self._num_format(num_id, ilvl or 0)
        return num_id, ilvl or 0, fmt

    def _num_format(self, num_id, ilvl) -> str:
        try:
            numbering = self.d.part.numbering_part.element
        except Exception:
            return "bullet"
        num = numbering.xpath(f'w:num[@w:numId="{num_id}"]')
        if not num:
            return "bullet"
        abs_id = num[0].xpath("w:abstractNumId/@w:val")
        if not abs_id:
            return "bullet"
        lv = numbering.xpath(
            f'w:abstractNum[@w:abstractNumId="{abs_id[0]}"]/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val')
        return lv[0] if lv else "decimal"

    _QT_LIST = {
        "bullet": QTextListFormat.ListDisc,
        "decimal": QTextListFormat.ListDecimal,
        "lowerLetter": QTextListFormat.ListLowerAlpha,
        "upperLetter": QTextListFormat.ListUpperAlpha,
        "lowerRoman": QTextListFormat.ListLowerRoman,
        "upperRoman": QTextListFormat.ListUpperRoman,
    }

    # ---- Einfügen ----------------------------------------------------
    def _start_block(self, bfmt: QTextBlockFormat, cfmt: QTextCharFormat, cursor: QTextCursor, reuse: bool):
        if reuse:
            cursor.setBlockFormat(bfmt)
            cursor.setBlockCharFormat(cfmt)
        else:
            cursor.insertBlock(bfmt, cfmt)

    def add_paragraph(self, para: Paragraph, cursor: QTextCursor, reuse: bool, in_table=False):
        bfmt = self._block_format(para)
        self._block_zusatz(bfmt, para)
        base_cfmt = QTextCharFormat()
        base_cfmt.setFontFamilies([self._font_name(_NoRun(), para)])
        base_cfmt.setFontPointSize(self._font_size_pt(_NoRun(), para))
        sz = _xp(para._p, "w:pPr/w:rPr/w:sz/@w:val")                       # Schrift des Absatzzeichens: sie bestimmt die Höhe leerer Absätze
        if sz and sz[0].isdigit():
            base_cfmt.setFontPointSize(int(sz[0]) / 2.0)
        self._start_block(bfmt, base_cfmt, cursor, reuse)

        info = None if in_table else self._list_info(para)
        if info is not None:
            num_id, ilvl, kind = info
            lf = QTextListFormat()
            lf.setStyle(self._QT_LIST.get(kind, QTextListFormat.ListDisc))
            lf.setIndent(ilvl + 1)
            key = (num_id, ilvl)
            if self.last_list and self.last_list[0] == key:
                self.last_list[1].add(cursor.block())
            else:
                self.last_list = (key, cursor.createList(lf))
        else:
            self.last_list = None

        self.deferred_break = False
        had_content = False
        had_content |= self._laufe(para._p, para, cursor)
        if self.deferred_break:
            # Umbruch am Absatzende: er gilt für den nächsten Absatz
            self.deferred_break = False
            self.pending_break = True
            self.pending_col = getattr(self, "deferred_col", False)
            return not had_content        # leerer Absatz → Block wiederverwenden
        return False

    def _laufe(self, eltern, para, cursor: QTextCursor, href: Optional[str] = None) -> bool:
        """Kinder eines Absatzes der Reihe nach: Läufe, Links, nachverfolgte Änderungen, Kommentar- und Textmarkenbereiche."""
        had = False
        for el in eltern:
            tag = el.tag
            if tag == qn("w:r"):
                had |= self._add_run(Run(el, para), para, cursor, href=href)
            elif tag == qn("w:hyperlink"):
                h = Hyperlink(el, para)
                url = h.url or ("#" + h.fragment if h.fragment else "")
                had |= self._laufe(el, para, cursor, href=url)
            elif tag in (qn("w:ins"), qn("w:del")):
                alt = self._aend
                self._aend = ("ins" if tag == qn("w:ins") else "del", el.get(qn("w:author")) or "", el.get(qn("w:date")) or "")
                had |= self._laufe(el, para, cursor, href=href)
                self._aend = alt
            elif tag == qn("w:commentRangeStart"):
                self._komm_offen.append(int(el.get(qn("w:id"))))
            elif tag == qn("w:commentRangeEnd"):
                i = int(el.get(qn("w:id")))
                if i in self._komm_offen:
                    self._komm_offen.remove(i)
            elif tag == qn("w:bookmarkStart"):
                n = el.get(qn("w:name")) or ""
                if n and not n.startswith("_"):
                    self._lz_offen[el.get(qn("w:id"))] = n
                    self._lz_pos[el.get(qn("w:id"))] = cursor.position()
            elif tag == qn("w:bookmarkEnd"):
                i = el.get(qn("w:id"))
                n = self._lz_offen.pop(i, None)
                if n is not None and self._lz_pos.pop(i, None) == cursor.position():
                    self._lz_wartend.append(n)         # leere Textmarke: hängt am folgenden Text
            elif tag in (qn("w:sdt"),):
                inhalt = el.find(qn("w:sdtContent"))
                if inhalt is not None:
                    alt_form, self._form = self._form, self._form_meta(el) or self._form
                    had |= self._laufe(inhalt, para, cursor, href=href)
                    self._form = alt_form
            elif tag == qn("w:fldSimple"):
                self._feld = {"instr": el.get(qn("w:instr")) or "", "phase": "r"}
                had |= self._laufe(el, para, cursor, href=href)
                self._feld_ende()
            elif tag == qn("w:smartTag"):
                had |= self._laufe(el, para, cursor, href=href)
        return had

    FELDER_BEHALTEN = ("SEQ", "CITATION")

    def _feld_ende(self) -> None:
        f, self._feld = self._feld, None
        if not f:
            return
        teile = f["instr"].split()
        if teile and teile[0].upper() == "XE" and self._letzter:
            term = f["instr"].split('"')[1] if '"' in f["instr"] else " ".join(teile[1:])
            c = QTextCursor(self.doc)
            c.setPosition(self._letzter[0]); c.setPosition(self._letzter[1], QTextCursor.KeepAnchor)
            nf = QTextCharFormat(); nf.setProperty(P_INDEX, term)
            c.mergeCharFormat(nf)

    def _form_meta(self, sdt) -> str:
        """Inhaltssteuerelement → Formularfeld-Daten (JSON); eigene Felder tragen sie im Tag, fremde (Word) werden aus der Art abgeleitet."""
        from lunivo_writer.editor import formular as FO
        pr = sdt.find(qn("w:sdtPr"))
        if pr is None:
            return ""
        tag = pr.find(qn("w:tag"))
        t = tag.get(qn("w:val")) if tag is not None else ""
        if t and t.startswith("lunivo:"):
            m = FO.meta_lesen(t[len("lunivo:"):])
            if m:
                return FO.meta_text(m)
        al = pr.find(qn("w:alias"))
        name = al.get(qn("w:val")) if al is not None else ""
        leer = pr.find(qn("w:showingPlcHdr")) is not None
        if pr.find(qn("w:dropDownList")) is not None or pr.find(qn("w:comboBox")) is not None:
            liste = pr.find(qn("w:dropDownList"))
            if liste is None:
                liste = pr.find(qn("w:comboBox"))
            opt = [li.get(qn("w:displayText")) or li.get(qn("w:value")) or "" for li in liste.findall(qn("w:listItem"))]
            m = FO.meta("auswahl", name, optionen=[o for o in opt if o])
        elif pr.find(qn("w:date")) is not None:
            m = FO.meta("datum", name)
        elif any(el.tag.endswith("}checkbox") for el in pr):
            m = FO.meta("kasten", name)
        elif pr.find(qn("w:text")) is not None:
            m = FO.meta("text", name)
        else:
            return ""
        if leer:
            m["leer"] = True
        return FO.meta_text(m)

    def _zusatz_formate(self, cfmt: QTextCharFormat) -> None:
        if self._feld and self._feld["phase"] == "r":
            teile = self._feld["instr"].split()
            if teile and teile[0].upper() in self.FELDER_BEHALTEN:
                cfmt.setProperty(P_FELD, self._feld["instr"].strip())
                if teile[0].upper() == "CITATION" and len(teile) > 1:
                    cfmt.setProperty(P_ZITAT, teile[1])
        if self._komm_offen:
            cfmt.setProperty(P_KOMMENTAR, self._komm_offen[-1])
        if self._form:
            cfmt.setProperty(P_FORM, self._form)
        if self._aend:
            cfmt.setProperty(P_AENDERUNG, self._aend[0])
            cfmt.setProperty(P_AUTOR, self._aend[1])
            cfmt.setProperty(P_DATUM, self._aend[2])
        namen = list(self._lz_offen.values()) + self._lz_wartend
        if namen and not cfmt.isAnchor():
            cfmt.setAnchor(True)
            cfmt.setAnchorNames(namen)

    _TOC_STILE = {"TOC Heading": "kopf", "TOC 1": "zeile1", "TOC 2": "zeile2", "TOC 3": "zeile3", "Table of Figures": "zeile1",
                  "Index 1": "zeile1", "Bibliography": "zeile1"}

    def _block_zusatz(self, bfmt: QTextBlockFormat, para: Paragraph) -> None:
        """Verzeichnis- und Beschriftungsabsätze an ihrer Formatvorlage wiedererkennen; Absatzrahmen aus w:pBdr."""
        ppr = para._p.pPr
        bdr = ppr.find(qn("w:pBdr")) if ppr is not None else None
        if bdr is not None:
            da = [(kurz, bdr.find(qn(f"w:{x}"))) for x, kurz in (("top", "t"), ("left", "l"), ("bottom", "b"), ("right", "r"), ("between", "h"))]
            da = [(k, e) for k, e in da if e is not None and e.get(qn("w:val"), "nil") not in ("nil", "none")]
            if da:
                e = da[0][1]
                col = e.get(qn("w:color")) or "000000"
                art = art_aus_docx(e.get(qn("w:val")))
                li = Linie(f"#{col.lower() if col.lower() != 'auto' else '000000'}", int(e.get(qn("w:sz")) or 8) / 8, art).mit_seiten("".join(k for k, _e in da))
                abst = {k: e2.get(qn("w:space")) for k, e2 in da}
                if any(v not in (None, "4") for v in abst.values()):          # eigener Abstand zum Text (w:space in pt)
                    li.abstand = ",".join(str(int(abst.get(k) or 4)) for k in "tblr")
                bfmt.setProperty(B_RAHMEN, li.text())
        try:
            name = self._ps(para).name
        except Exception:
            return
        if name in self._TOC_STILE:
            art = {"TOC 1": "toc", "TOC 2": "toc", "TOC 3": "toc", "Index 1": "index", "Bibliography": "lit"}.get(name)
            if name == "Table of Figures":
                art = "tab" if para.text.strip().lower().startswith(("tabelle", "table")) else "abb"
            elif name == "TOC Heading":
                nxt = para._p.getnext()
                art = "toc"
                if nxt is not None and nxt.tag == qn("w:p"):
                    nn = Paragraph(nxt, para._parent).style.name
                    art = {"Table of Figures": "abb", "Index 1": "index", "Bibliography": "lit"}.get(nn, "toc")
                    if nn == "Table of Figures" and Paragraph(nxt, para._parent).text.strip().lower().startswith(("tabelle", "table")):
                        art = "tab"
            bfmt.setProperty(B_TOC, self._TOC_STILE[name])
            bfmt.setProperty(B_ART, art or "toc")
        elif name == "Caption":
            erstes = para.text.strip().split(" ")[0] if para.text.strip() else ""
            bfmt.setProperty(B_CAP, erstes if erstes in BESCHRIFTUNGEN else (erstes or "Abbildung"))

    def _flush_break(self, cursor: QTextCursor):
        if self.deferred_break:
            self.deferred_break = False
            cursor.insertBlock()
            bf = cursor.blockFormat()
            bf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
            if getattr(self, "deferred_col", False):
                bf.setProperty(B_SPALTE, 1)
                self.deferred_col = False
            cursor.setBlockFormat(bf)

    def _add_run(self, run, para, cursor: QTextCursor, href: Optional[str] = None) -> bool:
        """→ True, wenn Text oder Bild eingefügt wurde."""
        added = False
        cfmt = self._char_format(run, para)
        self._zusatz_formate(cfmt)
        if href:
            cfmt.setAnchor(True)
            cfmt.setAnchorHref(href)
            cfmt.setForeground(QBrush(QColor("#0563C1")))
            cfmt.setFontUnderline(True)
        for child in run._r:
            tag = child.tag
            if tag == qn("w:fldChar"):
                typ = child.get(qn("w:fldCharType"))
                if typ == "begin":
                    self._feld = {"instr": "", "phase": "i"}
                elif typ == "separate" and self._feld:
                    self._feld["phase"] = "r"
                    cfmt = self._char_format(run, para)
                    self._zusatz_formate(cfmt)
                elif typ == "end":
                    self._feld_ende()
                continue
            if tag == qn("w:instrText"):
                if self._feld and self._feld["phase"] == "i":
                    self._feld["instr"] += child.text or ""
                continue
            if self._feld and self._feld["phase"] == "i":
                continue
            if tag == qn("w:t") or tag == qn("w:delText"):
                if child.text:
                    self._flush_break(cursor)
                    von = cursor.position()
                    cursor.insertText(child.text, cfmt)
                    self._letzter = (von, cursor.position())
                    self._lz_wartend = []
                    added = True
            elif tag == qn("w:tab"):
                self._flush_break(cursor)
                cursor.insertText("\t", cfmt)
                added = True
            elif tag in (qn("w:br"), qn("w:cr")):
                if child.get(qn("w:type")) in ("page", "column"):
                    self.deferred_break = True
                    self.deferred_col = child.get(qn("w:type")) == "column"
                else:
                    self._flush_break(cursor)
                    cursor.insertText("\u2028", cfmt)
            elif tag == qn("w:noBreakHyphen"):
                self._flush_break(cursor)
                cursor.insertText("\u2011", cfmt)
                added = True
            elif tag == qn("w:drawing"):
                self._flush_break(cursor)
                self._add_image(child, cursor, cfmt)
                added = True
            elif tag == _MC_ALT:                                      # Word legt Formen (Foto im Kreis) in „Alternativer Inhalt“: die Wahl nehmen
                wahl = child.find(_MC_CHOICE)
                for dr in (wahl.iter(qn("w:drawing")) if wahl is not None else []):
                    self._flush_break(cursor)
                    self._add_image(dr, cursor, cfmt)
                    added = True
            elif tag in (qn("w:footnoteReference"), qn("w:endnoteReference")):
                self._flush_break(cursor)
                art = "f" if tag == qn("w:footnoteReference") else "e"
                nid = int(child.get(qn("w:id")))
                self._note_n[art] += 1
                nf = QTextCharFormat(cfmt)
                nf.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
                nf.setProperty(P_NOTE, f"{art}:{nid}")
                cursor.insertText(self._note_marke(art, self._note_n[art]), nf)
                added = True
        return added

    @staticmethod
    def _note_marke(art: str, n: int) -> str:
        from lunivo_writer.editor.zusatz import roemisch
        return str(n) if art == "f" else roemisch(n)

    def _frei(self, drawing, teil=None, inline_ok: bool = False):
        """Ein Bild, das frei auf dem Blatt liegt (wp:anchor ohne Textumbruch) → (Angabe, QImage); ``False``, wenn es nicht darstellbar ist
        (Vektorbild); ``None``, wenn es kein solches Bild ist (steht im Text oder umflossen). Siehe editor/freie_bilder.py."""
        from lunivo_writer.editor.freie_bilder import emu_zu_px
        anker = drawing.xpath("wp:anchor") or (drawing.xpath("wp:inline") if inline_ok else [])
        if not anker or _xp(anker[0], "wp:wrapSquare | wp:wrapTight | wp:wrapThrough | wp:wrapTopAndBottom"):
            return None
        a = anker[0]
        blips = drawing.xpath(".//a:blip")
        if not blips:
            return None
        rid = blips[0].get(qn("r:embed"))
        try:
            blob = (teil or self.d.part).related_parts[rid].blob
        except KeyError:
            self.report.warn("Ein Bild fehlt im Dokument und wurde übersprungen.")
            return False
        from lunivo_writer.editor.freie_bilder import bild_aus_daten
        img = bild_aus_daten(blob)
        if img.isNull():
            self.report.warn("Ein Bild hat ein nicht lesbares Format (zum Beispiel eine Vektorzeichnung) und wird nicht angezeigt.")
            return False
        duo = _xp(drawing, ".//a:blip/a:duotone")
        if duo and len(duo[0]) >= 2:                                   # zweifarbig umgefärbt (zum Beispiel grau statt grün)
            from lunivo_writer.editor.freie_bilder import duoton, farbe_aus_element
            try:
                img = duoton(img, farbe_aus_element(duo[0][0], _NS), farbe_aus_element(duo[0][1], _NS))
            except Exception:                                          # noqa: BLE001 — ein ungewöhnlicher Effekt darf das Laden nicht verhindern
                pass
        ext = _xp(a, "wp:extent")
        b = emu_zu_px(ext[0].get("cx")) if ext else float(img.width())
        h = emu_zu_px(ext[0].get("cy")) if ext else float(img.height())
        rand = None
        for ln in _xp(drawing, ".//a:ln[a:solidFill]"):                  # Umrandung der Form (der weiße Ring um ein Foto)
            try:
                from lunivo_writer.editor.freie_bilder import farbe_aus_element
                f = farbe_aus_element(_xp(ln, "a:solidFill/*")[0], _NS)
                rand = {"b": max(0.5, emu_zu_px(ln.get("w") or 9525)), "farbe": f.name()}
                dash = _xp(ln, "a:prstDash/@val")
                rand["art"] = {"sysDot": "gepunktet", "dot": "gepunktet", "sysDash": "fein_gestrichelt", "dash": "gestrichelt", "lgDash": "gestrichelt",
                               "dashDot": "strichpunkt", "lgDashDot": "strichpunkt", "lgDashDotDot": "strichpunktpunkt",
                               "sysDashDot": "strichpunkt", "sysDashDotDot": "strichpunktpunkt"}.get(dash[0] if dash else "", "einfach")
            except Exception:                                              # noqa: BLE001
                rand = None
            break
        angabe = {"b": b, "h": h, "hinter": (a.get("behindDoc") or "0").lower() in ("1", "true"), "bh": "spalte", "bv": "absatz",
                  "xh": 0.0, "yv": 0.0, "ah": "", "av": "",
                  "ellipse": bool(_xp(drawing, ".//a:prstGeom[@prst='ellipse']"))}
        if rand:
            angabe["rand"] = rand
        fl = _xp(drawing, ".//pic:spPr/a:solidFill/*")  # Füllung der Form hinter dem Bild
        if fl:
            try:
                from lunivo_writer.editor.freie_bilder import farbe_aus_element
                angabe["fuell"] = {"art": "farbe", "farbe": farbe_aus_element(fl[0], _NS).name()}
            except Exception:                                              # noqa: BLE001
                pass
        rot = _xp(drawing, ".//a:xfrm/@rot")
        if rot and rot[0].lstrip("-").isdigit() and int(rot[0]):
            angabe["dreh"] = int(rot[0]) / 60000.0
        for achse, el in (("h", _xp(a, "wp:positionH")), ("v", _xp(a, "wp:positionV"))):
            if not el:
                continue
            bezug = el[0].get("relativeFrom", "")
            if achse == "h":
                angabe["bh"] = {"page": "seite", "margin": "rand", "leftMargin": "rand", "insideMargin": "rand", "outsideMargin": "rand",
                                "rightMargin": "rand"}.get(bezug, "spalte")
            else:
                angabe["bv"] = {"page": "seite", "margin": "rand", "topMargin": "rand", "insideMargin": "rand", "outsideMargin": "rand",
                                "bottomMargin": "rand"}.get(bezug, "absatz")
            al = _xp(el[0], "wp:align/text()")
            off = _xp(el[0], "wp:posOffset/text()")
            if al:
                angabe["ah" if achse == "h" else "av"] = al[0]
            elif off:
                angabe["xh" if achse == "h" else "yv"] = emu_zu_px(off[0])
        return angabe, img

    def _add_image(self, drawing, cursor: QTextCursor, cfmt: QTextCharFormat):
        frei = self._frei(drawing)
        if frei is False:
            return
        if frei is not None:
            from lunivo_writer.editor import freie_bilder as FB
            angabe, bild = frei
            name = FB.neuer_name(self.doc)
            self.doc.addResource(QTextDocument.ImageResource, QUrl(name), bild)
            angabe["n"] = name
            FB.an_absatz(cursor, angabe)
            return
        blips = drawing.xpath(".//a:blip")
        if not blips:
            return
        rid = blips[0].get(qn("r:embed"))
        try:
            blob = self.d.part.related_parts[rid].blob
        except KeyError:
            self.report.warn("Ein Bild fehlt im Dokument und wurde übersprungen.")
            return
        from lunivo_writer.editor.freie_bilder import bild_aus_daten
        img = bild_aus_daten(blob)
        if img.isNull():
            self.report.warn("Ein Bild hat ein nicht lesbares Format und wurde übersprungen.")
            return
        extent = drawing.xpath(".//wp:extent")
        w, h = float(img.width()), float(img.height())
        if extent:
            w = _emu_to_px(extent[0].get("cx"))
            h = _emu_to_px(extent[0].get("cy"))
        self.img_n += 1
        name = f"docx-bild-{self.img_n}"
        self.doc.addResource(QTextDocument.ImageResource, QUrl(name), img)
        imf = QTextImageFormat()
        imf.setName(name)
        imf.setWidth(w)
        imf.setHeight(h)
        alt = drawing.xpath(".//wp:docPr/@descr")
        if alt:
            imf.setProperty(STYLE_PROP + 1, alt[0])
        titel = drawing.xpath(".//wp:docPr/@title")
        if titel and titel[0].startswith("lunivo:"):
            imf.setProperty(P_OBJEKT, titel[0][len("lunivo:"):])
        seite = self._umbruch_seite(drawing, w)
        if seite:
            from lunivo_writer.editor.umbruch import absatz_nach_rahmen, schwebend_einfuegen
            bild_pos = schwebend_einfuegen(self.doc, cursor.block().position(), imf, seite)   # gehört zum Absatz: davor eingefügt
            text = absatz_nach_rahmen(self.doc, bild_pos)
            cursor.setPosition(text.position() + text.length() - 1)                          # weiter am Ende dieses Absatzes schreiben
            return
        cursor.insertImage(imf)

    def _umbruch_seite(self, drawing, breite_px: float) -> str:
        """„links“/„rechts“, wenn das Bild den Text umfließt (wrapSquare/Tight/Through); sonst leer."""
        anker = drawing.xpath("wp:anchor")
        if not anker or not _xp(anker[0], "wp:wrapSquare | wp:wrapTight | wp:wrapThrough"):
            return ""
        h = _xp(anker[0], "wp:positionH")
        if h:
            al = _xp(h[0], "wp:align/text()")
            if al:
                return "rechts" if al[0] in ("right", "outside") else "links"
            off = _xp(h[0], "wp:posOffset/text()")
            if off:
                try:
                    sec = self.d.sections[0]
                    inhalt = _emu_to_px(sec.page_width - sec.left_margin - sec.right_margin)
                except Exception:                                                      # noqa: BLE001
                    inhalt = 600.0
                return "rechts" if _emu_to_px(off[0]) + breite_px / 2 > inhalt / 2 else "links"
        return "links"

    EIGENE_MAX = 20                                                  # so viele Vorlagen nimmt ein Dokument höchstens in den Katalog auf

    def _eigene_vorlage(self, st, b) -> None:
        """Eine selbst angelegte Absatzvorlage des Dokuments: der Absatz behält ihren Namen; fehlt sie im Katalog, wird sie übernommen."""
        from docx.enum.style import WD_STYLE_TYPE
        from lunivo_writer.editor import stile
        if st.type != WD_STYLE_TYPE.PARAGRAPH or not st.name or "|" in st.name:
            return
        if not stile.ist_eigen(st.name):
            if st.name in stile.VORGABE or len(stile._NEU) >= self.EIGENE_MAX:
                return
            f, pf = st.font, st.paragraph_format
            try:
                farbe = ("#" + str(f.color.rgb)) if f.color is not None and f.color.type is not None and f.color.rgb is not None else ""
            except Exception:
                farbe = ""
            ausr = {WD_ALIGN_PARAGRAPH.LEFT: "links", WD_ALIGN_PARAGRAPH.CENTER: "mitte", WD_ALIGN_PARAGRAPH.RIGHT: "rechts",
                    WD_ALIGN_PARAGRAPH.JUSTIFY: "block"}.get(pf.alignment, "")
            stile.anlegen(st.name, groesse=f.size.pt if f.size else 0, fett=bool(f.bold), kursiv=bool(f.italic), farbe=farbe,
                          schrift=f.name or "", vor=pf.space_before.pt if pf.space_before is not None else 0,
                          nach=pf.space_after.pt if pf.space_after is not None else 0, ausr=ausr,
                          zeilen=float(pf.line_spacing) if isinstance(pf.line_spacing, float) else 0,
                          einzug=round(pf.left_indent.mm, 1) if pf.left_indent is not None else 0)
            self.neue_vorlagen.append(st.name)
        if stile.ist_eigen(st.name):
            b.setProperty(STYLE_PROP, stile.PRAEFIX + st.name)

    # ---- Tabellen ----------------------------------------------------
    @staticmethod
    def _zelle_format(qcell, cell) -> None:
        """Zellenschattierung (w:shd) und senkrechte Ausrichtung (w:vAlign) aus der DOCX-Zelle."""
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT as V
        cf = qcell.format().toTableCellFormat()
        geaendert = False
        shd = cell._tc.xpath("w:tcPr/w:shd")
        if shd and shd[0].get(qn("w:fill")) not in (None, "auto"):
            cf.setBackground(QBrush(QColor("#" + shd[0].get(qn("w:fill")))))
            geaendert = True
        va = cell.vertical_alignment
        if va in (V.CENTER, V.BOTTOM):
            cf.setVerticalAlignment(QTextCharFormat.AlignMiddle if va == V.CENTER else QTextCharFormat.AlignBottom)
            geaendert = True
        if geaendert:
            qcell.setFormat(cf)

    def add_table(self, table: Table):
        rows, cols = len(table.rows), len(table.columns)
        if rows == 0 or cols == 0:
            return
        tf = QTextTableFormat()
        tf.setCellPadding(4)
        tf.setCellSpacing(0)
        tf.setWidth(QTextLength(QTextLength.PercentageLength, 100))
        grid = table._tbl.xpath("w:tblPr/w:tblBorders") or (
            [1] if table.style is not None and "Grid" in (table.style.name or "") else [])
        tf.setBorder(0.5 if grid else 0)
        tf.setBorderCollapse(True)
        qt = self.cur.insertTable(rows, cols, tf)

        # Zellen mit Rechteck: gleiche tc → eine Zelle
        placed: dict[int, list] = {}
        for r in range(rows):
            for c in range(cols):
                tc = table.cell(r, c)._tc
                e = placed.setdefault(id(tc), [r, c, r, c, table.cell(r, c)])
                e[2] = max(e[2], r)
                e[3] = max(e[3], c)
        for r0, c0, r1, c1, cell in placed.values():
            if r1 > r0 or c1 > c0:
                qt.mergeCells(r0, c0, r1 - r0 + 1, c1 - c0 + 1)
        for r0, c0, _r1, _c1, cell in placed.values():
            ccur = qt.cellAt(r0, c0).firstCursorPosition()
            first = True
            for item in cell.iter_inner_content():
                if isinstance(item, Paragraph):
                    self.add_paragraph(item, ccur, reuse=first, in_table=True)
                    first = False
                else:
                    self.report.warn("Verschachtelte Tabellen werden nicht dargestellt.")
        for r0, c0, _r1, _c1, cell in placed.values():          # erst nach dem Inhalt: das Einfügen von Text setzt das Zellenformat zurück
            self._zelle_format(qt.cellAt(r0, c0), cell)
        self.cur = QTextCursor(self.doc)
        self.cur.movePosition(QTextCursor.End)
        self.fresh = True
        self.last_list = None

    # ---- Seite -------------------------------------------------------
    def read_page(self) -> PageSettings:
        """Einrichtung des ERSTEN Abschnitts; weitere Abschnitte stehen als Umbrüche im Text (siehe editor/abschnitte.py)."""
        return self._section_page(self.d.sections[0], erster=True)

    def _section_page(self, s, erster: bool = False, basis: PageSettings | None = None) -> PageSettings:
        ps = PageSettings()
        try:
            ps.width_mm = s.page_width.mm
            ps.height_mm = s.page_height.mm
            ps.margin_top_mm = s.top_margin.mm
            ps.margin_bottom_mm = s.bottom_margin.mm
            ps.margin_left_mm = s.left_margin.mm
            ps.margin_right_mm = s.right_margin.mm
        except Exception:
            self.report.warn("Seitengröße oder Ränder fehlen; es gilt A4 mit 2,5 cm Rand.")
        try:
            ps.kopf_abstand_mm, ps.fuss_abstand_mm = s.header_distance.mm, s.footer_distance.mm
        except Exception:                                                  # noqa: BLE001
            pass
        self._seitenschmuck(ps, s)
        if basis is not None:                                    # Hintergrund, Silbentrennung und Wasserzeichen gelten für das ganze Dokument
            ps.color, ps.silbentrennung, ps.watermark = basis.color, basis.silbentrennung, basis.watermark
        ps.kopf_verknuepft = (not erster) and s.header.is_linked_to_previous
        ps.fuss_verknuepft = (not erster) and s.footer.is_linked_to_previous
        ps.header = self._hf(s.header)
        ps.footer = self._hf(s.footer)
        ps.bilder = self._hf_bilder(s.header, "kopf", s, ps) + self._hf_bilder(s.footer, "fuss", s, ps)
        if s.different_first_page_header_footer:
            ps.diff_first = True
            ps.first_header = self._hf(s.first_page_header)
            ps.first_footer = self._hf(s.first_page_footer)
            ps.bilder += self._hf_bilder(s.first_page_header, "kopf1", s, ps) + self._hf_bilder(s.first_page_footer, "fuss1", s, ps)
        return ps

    def _hf_bilder(self, teil, ort: str, s, ps: PageSettings) -> list[dict]:
        """Frei positionierte Bilder einer Kopf- oder Fußzeile (zum Beispiel ein Band über dem Briefkopf). „Am Absatz“ wird auf das Blatt
        umgerechnet: der Absatz steht im Abstand der Kopfzeile vom oberen Rand, der der Fußzeile über dem unteren."""
        from lunivo_writer.models.page import mm_to_px
        if teil.is_linked_to_previous:
            return []
        out = []
        for dr in teil._element.iter(qn("w:drawing")):
            frei = self._frei(dr, teil.part, inline_ok=True)
            if not frei:
                continue
            angabe, bild = frei
            angabe["ort"], angabe["bild"] = ort, bild
            if dr.xpath("wp:inline"):                                  # Bild im Text der Kopf-/Fußzeile (Zierleiste): am Rand, am Absatz eingerückt
                ind = _xp(dr, "ancestor::w:p[1]/w:pPr/w:ind/@w:left")
                try:
                    angabe["xh"] = float(ind[0]) / 15.0 if ind else 0.0                      # Twips → px
                except ValueError:
                    angabe["xh"] = 0.0
                angabe["bh"] = "rand"
                try:
                    abstand = (s.header_distance if ort.startswith("kopf") else s.footer_distance).mm
                except Exception:                                                  # noqa: BLE001
                    abstand = 12.7
                angabe["bv"] = "seite"
                angabe["yv"] = mm_to_px(abstand) if ort.startswith("kopf") else mm_to_px(ps.height_mm - abstand) - angabe["h"]
                angabe["hinter"] = False
                out.append(angabe)
                continue
            if angabe["bv"] == "absatz":
                try:
                    abstand = (s.header_distance if ort.startswith("kopf") else s.footer_distance).mm
                except Exception:                                                  # noqa: BLE001
                    abstand = 12.7
                if ort.startswith("kopf"):
                    angabe["yv"] = mm_to_px(abstand) + angabe["yv"]
                else:
                    angabe["yv"] = mm_to_px(ps.height_mm - abstand) - 20 + angabe["yv"]
                angabe["bv"] = "seite"
            out.append(angabe)
        return out

    def _abschnitt_beginn(self, s) -> str:
        from lunivo_writer.editor.abschnitte import VOM_DOCX
        t = s._sectPr.find(qn("w:type"))
        return VOM_DOCX.get(t.get(qn("w:val")) if t is not None else "nextPage", "naechste")

    def _seitenschmuck(self, ps: PageSettings, sect) -> None:
        bg = self.d.element.find(qn("w:background"))
        if bg is not None:
            c = bg.get(qn("w:color"))
            if c and c.lower() not in ("auto", "ffffff"):
                ps.color = "#" + c.lower()
        if sect._sectPr.find(qn("w:lnNumType")) is not None:
            ps.zeilennummern = True
        co = sect._sectPr.find(qn("w:cols"))
        if co is not None:
            try:
                n = int(co.get(qn("w:num"), "1"))
            except ValueError:
                n = 1
            if n > 1:
                ps.spalten = min(n, 4)
                ps.spalten_abstand_mm = int(co.get(qn("w:space"), "720")) / 1440 * 25.4
                ps.spalten_linie = co.get(qn("w:sep"), "0").lower() in ("1", "true", "on")
        dt = self.d.settings.element.find(qn("w:defaultTabStop"))                      # Standardtabstopps des Dokuments (Twips)
        if dt is not None and dt.get(qn("w:val"), "").isdigit() and int(dt.get(qn("w:val"))) > 0:
            from PySide6.QtGui import QTextOption
            from lunivo_writer.models.page import mm_to_px
            opt = self.doc.defaultTextOption()
            opt.setTabStopDistance(mm_to_px(int(dt.get(qn("w:val"))) / 1440 * 25.4))
            self.doc.setDefaultTextOption(opt)
        ah = self.d.settings.element.find(qn("w:autoHyphenation"))
        if ah is not None and ah.get(qn("w:val"), "true").lower() not in ("false", "0", "off"):
            ps.silbentrennung = True
        pb = sect._sectPr.find(qn("w:pgBorders"))
        if pb is not None:
            e = next((x for x in (pb.find(qn("w:top")), pb.find(qn("w:left")), pb.find(qn("w:bottom")), pb.find(qn("w:right")))
                      if x is not None and x.get(qn("w:val"), "none") not in ("none", "nil")), None)
            if e is not None:
                col = (e.get(qn("w:color")) or "000000")
                da = "".join(k for x, k in (("top", "t"), ("bottom", "b"), ("left", "l"), ("right", "r"))
                             if pb.find(qn(f"w:{x}")) is not None and pb.find(qn(f"w:{x}")).get(qn("w:val"), "none") not in ("none", "nil"))
                ps.border = {"seiten": da or "tblr", "farbe": "#" + (col if col.lower() != "auto" else "000000").lower(),
                             "breite_pt": int(e.get(qn("w:sz")) or 8) / 8.0,
                             "art": art_aus_docx(e.get(qn("w:val"))),
                             "abstand_mm": int(e.get(qn("w:space")) or 24) / 72 * 25.4}
                from lunivo_writer.editor.seitenrahmen import KUNST
                wert = e.get(qn("w:val"))
                if wert in [k for k, _n in KUNST if k]:                              # Kunstrahmen
                    ps.border["kunst"] = wert
                if pb.get(qn("w:offsetFrom")) == "text":                              # Abstand vom Text, je Seite (pt)
                    ps.border["messen"] = "text"
                ab = [int((pb.find(qn(f"w:{x}")).get(qn("w:space")) if pb.find(qn(f"w:{x}")) is not None else None) or 0) for x in ("top", "bottom", "left", "right")]
                ps.border["abstand"] = ",".join(str(x) for x in ab)
                ps.border["vordergrund"] = pb.get(qn("w:zOrder"), "front") != "back"
                st = self.d.settings.element
                ps.border["kopf"] = st.find(qn("w:bordersDoNotSurroundHeader")) is None
                ps.border["fuss"] = st.find(qn("w:bordersDoNotSurroundFooter")) is None
                ps.border["ausrichten"] = st.find(qn("w:alignBordersAndEdges")) is not None

    def _hf(self, part) -> HeaderFooter:
        if part.is_linked_to_previous or not part.paragraphs:
            return HeaderFooter()
        lines, align = [], "left"
        for i, p in enumerate(part.paragraphs):
            txt = ""
            for el in p._p.iter():
                if el.tag == qn("w:t"):
                    txt += el.text or ""
                elif el.tag == qn("w:tab"):
                    txt += "\t"
                elif el.tag == qn("w:instrText") and "PAGE" in (el.text or "").upper() \
                        and "NUMPAGES" not in (el.text or "").upper():
                    txt += "{PAGE}"
                elif el.tag == qn("w:instrText") and "NUMPAGES" in (el.text or "").upper():
                    txt += "{PAGES}"
                elif el.tag == qn("w:fldSimple"):
                    instr = (el.get(qn("w:instr")) or "").upper()
                    txt += "{PAGES}" if "NUMPAGES" in instr else "{PAGE}" if "PAGE" in instr else ""
            if i == 0:
                a = self._ausrichtung(p, p.style)
                align = {Qt.AlignHCenter: "center", Qt.AlignRight: "right"}.get(a, "left")
            lines.append(txt)
        text = "\n".join(lines).rstrip("\n")
        return HeaderFooter(text, align)

    # ---- Hauptlauf ---------------------------------------------------
    def run(self):
        body = self.d.element.body
        for pat, msg in _UNSUPPORTED:
            try:
                if body.xpath(pat):
                    self.report.warn(msg)
            except Exception:
                pass
        def kinder(eltern):
            for c in eltern.iterchildren():
                if c.tag == qn("w:sdt"):
                    inhalt = c.find(qn("w:sdtContent"))
                    if inhalt is not None:
                        yield from kinder(inhalt)
                else:
                    yield c
        self._erste_seite = self.read_page()
        for child in kinder(body):
            if child.tag == qn("w:p"):
                self.fresh = self.add_paragraph(Paragraph(child, self.d._body), self.cur, reuse=self.fresh)
            elif child.tag == qn("w:tbl"):
                self.add_table(Table(child, self.d._body))
        # leeren Schlussblock nach Tabelle behalten (Qt braucht ihn)
        page = self._erste_seite
        df = QFont(self.def_font or "Calibri")
        df.setPointSizeF(self.def_size or 11.0)
        self.doc.setDefaultFont(df)
        props = self.d.core_properties
        meta = {"title": props.title or "", "author": props.author or ""}
        meta.update(self.zus)
        lang = self.d.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr") + "/" + qn("w:lang"))
        if lang is not None and lang.get(qn("w:val")):
            meta["sprache"] = lang.get(qn("w:val"))
        self.doc.setModified(False)
        return self.doc, page, meta, self.report


class _NoRun:
    """Platzhalter, um Absatz-Standardwerte ohne Run aufzulösen."""
    class _F:
        name = None
        size = None
        bold = italic = underline = strike = superscript = subscript = None
        all_caps = small_caps = outline = hidden = None
        color = None
    font = _F()
    style = None


_ERSATZ = {"Calibri Light": ["Carlito"], "Calibri": ["Carlito"], "Cambria": ["Caladea"], "Segoe UI": ["Noto Sans"]}


def _schriften_ersatz() -> None:
    """Word-Schriften ohne Entsprechung im System ersetzt Qt sonst durch irgendeine breite Schrift (so wurde aus „Calibri Light“ Noto Sans).
    Die metrikgleichen Schriften (Carlito, Caladea) halten Zeilenhöhe und Umbruch gleich — wie es LibreOffice mit seiner Ersetzungstabelle tut."""
    from PySide6.QtGui import QFont, QFontDatabase
    vorhanden = set(QFontDatabase.families())
    for name, ersatz in _ERSATZ.items():
        wahl = [e for e in ersatz if e in vorhanden]
        if wahl and name not in vorhanden:
            QFont.insertSubstitutions(name, wahl)


def read_docx(path: str):
    """→ (QTextDocument, PageSettings, meta-dict, IOReport). Wirft DocumentIOError."""
    _schriften_ersatz()
    return _Reader(path).run()
